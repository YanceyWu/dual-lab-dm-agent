from __future__ import annotations

import json
import sqlite3
import threading
from datetime import datetime, timedelta, timezone
from pathlib import Path

import pytest
from openpyxl import Workbook
from typer.testing import CliRunner

from pm_agent.cli.app import app
from pm_agent.config import settings
from pm_agent.database import repository as legacy_repository
from pm_agent.database.bootstrap import main as init_db
from pm_agent.data_onboarding import service as data_onboarding_service
from pm_agent.resource_intelligence.read_model import get_effective_capacity
from pm_agent.use_cases import hiref_management_service
from pm_agent.workbook_onboarding.parser import (
    EXPECTED_HEADERS,
    WorkbookParseError,
    parse_workbook,
)
from pm_agent.workbook_onboarding.presets import (
    DEFAULT_WORKBOOK_MAPPING_PRESET_ID,
    WORKBOOK_MAPPING_PRESET_ID_WITH_ALIASES,
    get_workbook_preset,
    list_workbook_presets,
    serialize_workbook_preset,
    validate_preset_registry,
)
from pm_agent.workbook_onboarding.service import (
    export_current_state_workbook,
    import_workbook,
    preview_workbook_import,
)
from pm_agent.workbook_onboarding import service as workbook_service
from pm_agent.workbook_onboarding.hiref_bridge import export_snapshot as hiref_export_snapshot
from pm_agent.workbook_onboarding.repository import (
    advance_profile_revision,
    resolve_plan_identity,
)
from pm_agent.workbook_onboarding.validator import validate_workbook
from pm_agent.workforce_planning_import import repository as workforce_repository
from pm_agent.workforce_planning_import.read_model import source_export_snapshot

runner = CliRunner()


def _onboard_workbook(
    path: Path, *, db_path: Path, profile_key: str | None = None
) -> dict[str, object]:
    """Publish a workbook through the supported profile/preview/confirm path."""
    profile_key = profile_key or f"export-{path.stem}"
    saved = data_onboarding_service.save_source_profile(
        profile_key=profile_key,
        source_type="workbook",
        source_locator=str(path),
        db_path=db_path,
    )
    assert saved["status"] == "saved"
    preview = data_onboarding_service.preview_source_profile(profile_key, db_path=db_path)
    assert preview["status"] == "previewed", preview
    confirmed = data_onboarding_service.confirm_onboarding_run(
        str(preview["run_id"]), db_path=db_path
    )
    assert confirmed["status"] == "completed", confirmed
    return confirmed


def _write_workbook(
    path: Path,
    *,
    setup_rows: list[list[object]],
    member_rows: list[list[object]],
    project_rows: list[list[object]],
    allocation_rows: list[list[object]],
    capacity_rows: list[list[object]],
    sheet_names: tuple[str, str, str, str, str] = (
        "Setup",
        "Members",
        "Projects",
        "Allocations",
        "Capacity",
    ),
    header_overrides: dict[str, list[str]] | None = None,
    extra_sheets: dict[str, list[list[object]]] | None = None,
) -> Path:
    workbook = Workbook()
    first = workbook.active
    first.title = sheet_names[0]
    for title in sheet_names[1:]:
        workbook.create_sheet(title)
    header_overrides = header_overrides or {}
    extra_sheets = extra_sheets or {}
    logical_sheet_names = ("Setup", "Members", "Projects", "Allocations", "Capacity")
    rows_by_sheet = {
        "Setup": setup_rows,
        "Members": member_rows,
        "Projects": project_rows,
        "Allocations": allocation_rows,
        "Capacity": capacity_rows,
    }
    for logical_name, worksheet in zip(logical_sheet_names, workbook.worksheets, strict=True):
        worksheet.append(header_overrides.get(logical_name, EXPECTED_HEADERS[logical_name]))
        for row in rows_by_sheet[logical_name]:
            worksheet.append(row)
    for logical_name, rows in extra_sheets.items():
        worksheet = workbook.create_sheet(logical_name)
        worksheet.append(header_overrides.get(logical_name, EXPECTED_HEADERS[logical_name]))
        for row in rows:
            worksheet.append(row)
    workbook.save(path)
    return path


def _baseline_workbook(
    path: Path,
    *,
    allocation_value: float = 0.5,
    include_capacity: bool = False,
    capacity_rows: list[list[object]] | None = None,
    members: list[list[object]] | None = None,
    projects: list[list[object]] | None = None,
    allocations: list[list[object]] | None = None,
    start_month: str = "2026-09",
    end_month: str = "2026-09",
    sheet_names: tuple[str, str, str, str, str] = (
        "Setup",
        "Members",
        "Projects",
        "Allocations",
        "Capacity",
    ),
    header_overrides: dict[str, list[str]] | None = None,
    extra_sheets: dict[str, list[list[object]]] | None = None,
) -> Path:
    return _write_workbook(
        path,
        setup_rows=[["FY26 Q4 Baseline", "2026-08-15", start_month, end_month]],
        member_rows=members
        or [
            [
                "WD100001",
                "Alex Example",
                "LTFTE",
                "active",
                None,
                None,
                "Engineer",
                7,
                "2026-09-01",
                None,
            ]
        ],
        project_rows=projects
        or [["RP-PROJ-001", "Project Atlas", "active", 2, "2026-09-01", "2026-12-31"]],
        allocation_rows=allocations
        or [["WD100001", "RP-PROJ-001", "2026-09", allocation_value]],
        capacity_rows=(
            capacity_rows
            if capacity_rows is not None
            else ([["WD100001", "2026-09", 0.0, 0.0, 0.0]] if include_capacity else [])
        ),
        sheet_names=sheet_names,
        header_overrides=header_overrides,
        extra_sheets=extra_sheets,
    )


def _confirm_staffing_adjustment(
    *,
    project_id: str,
    member_id: str,
    plan_version_id: str,
    month: str,
    allocation: float,
) -> None:
    proposal_id = "proposal-workbook-adjustment-001"
    token = "token-workbook-adjustment-001"
    legacy_repository.create_staffing_proposal(
        {
            "proposal_id": proposal_id,
            "expires_at": (
                datetime.now(timezone.utc) + timedelta(days=1)
            ).isoformat(timespec="seconds"),
            "confirmation_token": token,
            "request": {
                "project_id": project_id,
                "start_period": month,
                "role": "Engineer",
            },
            "evidence": {},
            "proposal": {
                "plan_version_id": plan_version_id,
                "periods": [{"year": int(month[:4]), "month": int(month[5:7])}],
                "selections": [{"member_id": member_id, "allocation": allocation}],
                "candidates": [],
            },
        }
    )
    result = legacy_repository.confirm_staffing_proposal(
        proposal_id,
        token,
        {
            "plan_version_id": plan_version_id,
            "periods": [{"year": int(month[:4]), "month": int(month[5:7])}],
            "selections": [{"member_id": member_id, "allocation": allocation}],
            "candidates": [],
        },
    )
    assert result["status"] == "confirmed"


ALIAS_SHEET_NAMES = (
    "Plan Setup",
    "Team Members",
    "Project List",
    "Project Allocations",
    "Member Capacity",
)

ALIAS_HEADERS = {
    "Setup": ["plan_name", "snapshot_date", "start_month", "end_month"],
    "Members": [
        "member_key",
        "member_name",
        "fte_type",
        "status",
        "hiref_id",
        "hiref_end_date",
        "role",
        "level",
        "effective_start",
        "effective_until",
        "next_hiref",
    ],
    "Projects": [
        "project_id",
        "display_name",
        "status",
        "priority",
        "start_date",
        "target_end_date",
    ],
    "Allocations": [
        "member_key",
        "project_id",
        "month",
        "allocation_fraction",
        "open_hiref_id",
    ],
    "Capacity": [
        "member_key",
        "month",
        "leave_fraction",
        "bau_fraction",
        "non_project_allocation",
    ],
}


def test_workbook_preset_registry_lists_and_validates_packaged_presets() -> None:
    validate_preset_registry()
    presets = list_workbook_presets()

    assert [preset.mapping_preset_id for preset in presets] == [
        DEFAULT_WORKBOOK_MAPPING_PRESET_ID,
        WORKBOOK_MAPPING_PRESET_ID_WITH_ALIASES,
    ]

    alias_preset = serialize_workbook_preset(
        get_workbook_preset(WORKBOOK_MAPPING_PRESET_ID_WITH_ALIASES)
    )
    assert alias_preset["status"] == "active"
    setup_section = next(
        section for section in alias_preset["sections"] if section["section_key"] == "setup"
    )
    assert setup_section["accepted_sheet_names"] == ["Setup", "Plan Setup"]
    assert setup_section["header_aliases"]["plan_version_name"] == ["plan_name"]


def test_parse_workbook_accepts_alias_preset_layout_and_tracks_resolution(
    tmp_path: Path,
) -> None:
    workbook_path = _baseline_workbook(
        tmp_path / "alias-layout.xlsx",
        include_capacity=True,
        sheet_names=ALIAS_SHEET_NAMES,
        header_overrides=ALIAS_HEADERS,
    )

    parsed = parse_workbook(
        workbook_path,
        mapping_preset_id=WORKBOOK_MAPPING_PRESET_ID_WITH_ALIASES,
    )
    result = validate_workbook(parsed)

    assert result.workbook is not None
    warning_codes = {item.code for item in result.warnings}
    assert "WORKBOOK_SHEET_ALIAS_USED" in warning_codes
    assert "WORKBOOK_HEADER_ALIAS_USED" in warning_codes
    assert parsed.preset_resolution.mapping_preset_id == WORKBOOK_MAPPING_PRESET_ID_WITH_ALIASES
    assert parsed.preset_resolution.matched_via_alias is True
    sections = {
        section.section_key: section for section in parsed.preset_resolution.sections
    }
    assert sections["setup"].resolved_sheet_name == "Plan Setup"
    assert sections["setup"].matched_via_alias is True
    assert any(
        match.field_key == "plan_version_name"
        and match.resolved_header_name == "plan_name"
        and match.matched_via_alias is True
        for match in sections["setup"].header_matches
    )


def test_parse_workbook_rejects_ambiguous_alias_header_match(tmp_path: Path) -> None:
    workbook_path = _baseline_workbook(
        tmp_path / "ambiguous-alias.xlsx",
        sheet_names=ALIAS_SHEET_NAMES,
        header_overrides={
            **ALIAS_HEADERS,
            "Setup": ["plan_version_name", "plan_name", "start_month", "end_month"],
        },
    )

    with pytest.raises(
        WorkbookParseError, match="WORKBOOK_HEADER_AMBIGUOUS:Plan Setup:plan_version_name"
    ):
        parse_workbook(
            workbook_path,
            mapping_preset_id=WORKBOOK_MAPPING_PRESET_ID_WITH_ALIASES,
        )


def test_parse_workbook_preserves_alias_resolution_on_row_width_failure(
    tmp_path: Path,
) -> None:
    workbook_path = _write_workbook(
        tmp_path / "alias-row-width.xlsx",
        setup_rows=[["FY26 Q4 Baseline", "2026-08-15", "2026-09", "2026-09", "extra"]],
        member_rows=[],
        project_rows=[],
        allocation_rows=[],
        capacity_rows=[],
        header_overrides=ALIAS_HEADERS,
    )

    with pytest.raises(WorkbookParseError) as exc_info:
        parse_workbook(
            workbook_path,
            mapping_preset_id=WORKBOOK_MAPPING_PRESET_ID_WITH_ALIASES,
        )

    error = exc_info.value
    assert str(error) == "WORKBOOK_ROW_WIDTH_INVALID:Setup:2"
    assert error.preset_resolution is not None
    assert error.preset_resolution.matched_via_alias is True
    setup_section = next(
        section
        for section in error.preset_resolution.sections
        if section.section_key == "setup"
    )
    assert any(
        match.field_key == "plan_version_name"
        and match.resolved_header_name == "plan_name"
        and match.matched_via_alias is True
        for match in setup_section.header_matches
    )


def test_preview_workbook_explicit_default_preset_matches_implicit_default(
    isolated_db: Path, tmp_path: Path
) -> None:
    init_db(quiet=True)
    workbook_path = _baseline_workbook(
        tmp_path / "default-preset-regression.xlsx",
        include_capacity=True,
    )

    implicit = preview_workbook_import(workbook_path, db_path=isolated_db)
    explicit = preview_workbook_import(
        workbook_path,
        mapping_preset_id=DEFAULT_WORKBOOK_MAPPING_PRESET_ID,
        db_path=isolated_db,
    )

    assert implicit["status"] == "previewed"
    assert explicit["status"] == "previewed"
    assert implicit["counts"] == explicit["counts"]
    assert implicit["plan_version"] == explicit["plan_version"]
    assert implicit["workforce_package"] == explicit["workforce_package"]
    assert implicit["current_state_staffing_package"] == explicit["current_state_staffing_package"]
    assert implicit["capacity_package"] == explicit["capacity_package"]


def test_export_current_state_workbook_is_atomic_metadata_only_and_roundtrips(
    isolated_db: Path, tmp_path: Path
) -> None:
    init_db(quiet=True)
    source = _baseline_workbook(
        tmp_path / "export-source.xlsx", include_capacity=True
    )
    imported = _onboard_workbook(source, db_path=isolated_db)
    assert imported["status"] == "completed"

    output = tmp_path / "planning-export.xlsx"
    result = export_current_state_workbook(output, db_path=isolated_db)

    assert result["status"] == "success"
    assert result["output_path"] == str(output.resolve())
    assert result["sha256"]
    assert result["counts"]["capacity"] == 1
    parsed = parse_workbook(output)
    assert parsed.setup_rows[0].plan_version_name == "FY26 Q4 Baseline"
    assert (parsed.setup_rows[0].start_month, parsed.setup_rows[0].end_month) == (
        "2026-09", "2026-09"
    )
    assert export_current_state_workbook(output, db_path=isolated_db)["warnings"] == [
        "WORKBOOK_EXPORT_OUTPUT_EXISTS"
    ]
    assert export_current_state_workbook(output, db_path=isolated_db, overwrite=True)[
        "status"
    ] == "success"


def test_export_rejects_racing_target_without_overwrite(
    isolated_db: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    init_db(quiet=True)
    assert _onboard_workbook(
        _baseline_workbook(tmp_path / "race-source.xlsx", include_capacity=True),
        db_path=isolated_db,
    )["status"] == "completed"
    output = tmp_path / "race.xlsx"

    def racing_link(source: Path, destination: Path) -> None:
        destination.write_bytes(b"racing-owner")
        raise FileExistsError

    monkeypatch.setattr(workbook_service.os, "link", racing_link)
    result = export_current_state_workbook(output, db_path=isolated_db)

    assert result == {"status": "failed", "warnings": ["WORKBOOK_EXPORT_OUTPUT_EXISTS"]}
    assert output.read_bytes() == b"racing-owner"


def test_export_uses_source_owner_snapshots_and_cli_emits_one_json_line(
    isolated_db: Path, tmp_path: Path
) -> None:
    init_db(quiet=True)
    assert _onboard_workbook(
        _baseline_workbook(tmp_path / "public-contract.xlsx", include_capacity=True),
        db_path=isolated_db,
    )["status"] == "completed"

    output = tmp_path / "cli-export.xlsx"
    result = runner.invoke(app, ["onboarding", "export-workbook", "--output", str(output)])
    assert result.exit_code == 0, result.output
    assert len(result.output.strip().splitlines()) == 1
    payload = json.loads(result.output)
    assert payload["status"] == "success"
    assert payload["evidence"]["plan_version_id"]
    assert payload["evidence"]["capacity_source"] is not None


def test_export_maps_legacy_member_level_labels_to_v1_levels(
    isolated_db: Path, tmp_path: Path
) -> None:
    init_db(quiet=True)
    assert _onboard_workbook(
        _baseline_workbook(tmp_path / "legacy-level.xlsx", include_capacity=True),
        db_path=isolated_db,
    )["status"] == "completed"
    with sqlite3.connect(isolated_db) as database:
        database.execute("UPDATE employees SET level='senior' WHERE id='WD100001'")
    output = tmp_path / "legacy-level-export.xlsx"
    assert export_current_state_workbook(output, db_path=isolated_db)["status"] == "success"
    assert parse_workbook(output).members[0].level == 7


def test_export_rejects_stfte_source_without_current_hiref_before_writing(
    isolated_db: Path, tmp_path: Path
) -> None:
    init_db(quiet=True)
    assert _onboard_workbook(
        _baseline_workbook(tmp_path / "missing-hiref.xlsx", include_capacity=True),
        db_path=isolated_db,
    )["status"] == "completed"
    with sqlite3.connect(isolated_db) as database:
        database.execute(
            "UPDATE employees SET resource_type='STFTE',current_hiref='',billing_end_date='' WHERE id='WD100001'"
        )
    output = tmp_path / "missing-hiref-export.xlsx"
    assert export_current_state_workbook(output, db_path=isolated_db) == {
        "status": "failed", "warnings": ["WORKBOOK_EXPORT_STFTE_HIREF_REQUIRED"]
    }
    assert not output.exists()


def test_export_ignores_stale_or_mismatched_derived_publications(
    isolated_db: Path, tmp_path: Path
) -> None:
    init_db(quiet=True)
    assert _onboard_workbook(
        _baseline_workbook(tmp_path / "stale-export.xlsx", include_capacity=True),
        db_path=isolated_db,
    )["status"] == "completed"
    with sqlite3.connect(isolated_db) as database:
        database.execute("UPDATE resource_capacity_observations SET observed_at='2000-01-01T00:00:00+00:00'")
        database.execute("DELETE FROM current_state_staffing_members")
        database.execute("DELETE FROM contract_coverage_members")
        database.execute("UPDATE data_onboarding_publication_links SET domain_publication_id='wrong-publication' WHERE capability_key='contract_coverage'")
    result = export_current_state_workbook(tmp_path / "stale.xlsx", db_path=isolated_db)
    assert result["status"] == "success"


def test_export_uses_latest_workbook_capacity_generation_even_when_empty(
    isolated_db: Path, tmp_path: Path
) -> None:
    init_db(quiet=True)
    first = _baseline_workbook(tmp_path / "capacity-a.xlsx", include_capacity=True)
    assert _onboard_workbook(first, db_path=isolated_db)["status"] == "completed"
    second = _baseline_workbook(tmp_path / "capacity-b.xlsx", include_capacity=False)
    confirmed = _onboard_workbook(second, db_path=isolated_db)
    assert confirmed["status"] == "completed", confirmed
    with sqlite3.connect(isolated_db) as database:
        database.execute(
            "UPDATE data_onboarding_runs SET completed_at='2026-08-09T00:00:00+00:00', "
            "created_at='2026-08-09T00:00:00+00:00' WHERE source_type='workbook'"
        )
    exported = export_current_state_workbook(tmp_path / "capacity-b-export.xlsx", db_path=isolated_db)
    assert exported["status"] == "success"
    assert exported["counts"]["capacity"] == 0


def test_export_rejects_boolean_capacity_row_count(
    isolated_db: Path, tmp_path: Path
) -> None:
    init_db(quiet=True)
    assert _onboard_workbook(
        _baseline_workbook(tmp_path / "boolean-row-count.xlsx", include_capacity=True),
        db_path=isolated_db,
    )["status"] == "completed"
    with sqlite3.connect(isolated_db) as database:
        preview = json.loads(database.execute(
            "SELECT source_preview_json FROM data_onboarding_runs WHERE source_type='workbook'"
        ).fetchone()[0])
        preview["workbook_export_metadata"]["capacity"]["row_count"] = True
        database.execute(
            "UPDATE data_onboarding_runs SET source_preview_json=?", [json.dumps(preview)]
        )
    assert export_current_state_workbook(
        tmp_path / "boolean-row-count-export.xlsx", db_path=isolated_db
    ) == {"status": "failed", "warnings": ["WORKBOOK_EXPORT_SOURCE_CONTEXT_INVALID"]}


def test_export_fails_closed_when_current_capacity_source_is_missing(
    isolated_db: Path, tmp_path: Path
) -> None:
    init_db(quiet=True)
    assert _onboard_workbook(
        _baseline_workbook(tmp_path / "capacity-old.xlsx", include_capacity=True),
        db_path=isolated_db, profile_key="capacity-currentness",
    )["status"] == "completed"
    assert _onboard_workbook(
        _baseline_workbook(tmp_path / "capacity-current.xlsx", include_capacity=True),
        db_path=isolated_db, profile_key="capacity-currentness",
    )["status"] == "completed"
    with sqlite3.connect(isolated_db) as database:
        preview = json.loads(database.execute(
            """SELECT source_preview_json FROM data_onboarding_runs
               WHERE source_type='workbook' AND status='completed'
               ORDER BY completed_at DESC,created_at DESC,rowid DESC LIMIT 1"""
        ).fetchone()[0])
        package_id = preview["workbook_export_metadata"]["capacity"]["package_id"]
        database.execute(
            "DELETE FROM resource_capacity_import_sessions WHERE package_id=?", [package_id]
        )
    output = tmp_path / "missing-capacity-source.xlsx"
    assert export_current_state_workbook(output, db_path=isolated_db) == {
        "status": "failed", "warnings": ["WORKBOOK_EXPORT_CAPACITY_SOURCE_MISSING"]
    }
    assert not output.exists()


def test_export_fails_closed_when_current_capacity_source_is_corrupt(
    isolated_db: Path, tmp_path: Path
) -> None:
    init_db(quiet=True)
    assert _onboard_workbook(
        _baseline_workbook(tmp_path / "capacity-corrupt.xlsx", include_capacity=True),
        db_path=isolated_db,
    )["status"] == "completed"
    with sqlite3.connect(isolated_db) as database:
        database.execute(
            "UPDATE resource_capacity_import_sessions SET package_json='{}' WHERE status='completed'"
        )
    output = tmp_path / "corrupt-capacity-source.xlsx"
    assert export_current_state_workbook(output, db_path=isolated_db) == {
        "status": "failed", "warnings": ["WORKBOOK_EXPORT_CAPACITY_SOURCE_INVALID"]
    }
    assert not output.exists()


def test_export_rejects_semantically_valid_tampered_capacity_package(
    isolated_db: Path, tmp_path: Path
) -> None:
    init_db(quiet=True)
    assert _onboard_workbook(
        _baseline_workbook(tmp_path / "capacity-fingerprint.xlsx", include_capacity=True),
        db_path=isolated_db,
    )["status"] == "completed"
    with sqlite3.connect(isolated_db) as database:
        raw = database.execute(
            "SELECT package_json FROM resource_capacity_import_sessions WHERE status='completed'"
        ).fetchone()[0]
        package = json.loads(raw)
        package["observations"][0]["fraction"] = 0.25
        database.execute(
            "UPDATE resource_capacity_import_sessions SET package_json=? WHERE status='completed'",
            [json.dumps(package)],
        )
    assert export_current_state_workbook(
        tmp_path / "tampered-capacity.xlsx", db_path=isolated_db
    ) == {"status": "failed", "warnings": ["WORKBOOK_EXPORT_CAPACITY_SOURCE_INVALID"]}


def test_export_preserves_authoritative_setup_horizon_with_sparse_facts(
    isolated_db: Path, tmp_path: Path
) -> None:
    init_db(quiet=True)
    source = _baseline_workbook(
        tmp_path / "sparse-horizon.xlsx", include_capacity=False,
        start_month="2026-09", end_month="2026-12",
        allocations=[["WD100001", "RP-PROJ-001", "2026-09", 0.5]],
    )
    assert _onboard_workbook(source, db_path=isolated_db)["status"] == "completed"
    output = tmp_path / "sparse-horizon-export.xlsx"
    assert export_current_state_workbook(output, db_path=isolated_db)["status"] == "success"
    parsed = parse_workbook(output)
    assert (parsed.setup_rows[0].start_month, parsed.setup_rows[0].end_month) == (
        "2026-09", "2026-12"
    )


def test_export_rejects_member_status_not_representable_by_v1(
    isolated_db: Path, tmp_path: Path
) -> None:
    init_db(quiet=True)
    assert _onboard_workbook(
        _baseline_workbook(tmp_path / "member-status.xlsx", include_capacity=True),
        db_path=isolated_db,
    )["status"] == "completed"
    with sqlite3.connect(isolated_db) as database:
        database.execute("UPDATE employees SET status='on_leave' WHERE id='WD100001'")
    output = tmp_path / "member-status-export.xlsx"
    assert export_current_state_workbook(output, db_path=isolated_db) == {
        "status": "failed", "warnings": ["WORKBOOK_EXPORT_MEMBER_STATUS_UNSUPPORTED"]
    }
    assert not output.exists()


def test_export_rejects_multiple_active_plans_without_selector(
    isolated_db: Path, tmp_path: Path
) -> None:
    init_db(quiet=True)
    assert _onboard_workbook(
        _baseline_workbook(tmp_path / "ambiguous-export.xlsx", include_capacity=True),
        db_path=isolated_db,
    )["status"] == "completed"
    with sqlite3.connect(isolated_db) as database:
        database.execute(
            """INSERT INTO plan_versions(plan_version_id,version_name,scenario_type,as_of_date,version_status)
               VALUES ('plan-other-active','Other active','baseline','2026-08-15','active')"""
        )
    assert export_current_state_workbook(tmp_path / "ambiguous.xlsx", db_path=isolated_db) == {
        "status": "failed", "warnings": ["WORKBOOK_EXPORT_ACTIVE_PLAN_AMBIGUOUS"]
    }
    selected = export_current_state_workbook(
        tmp_path / "out-of-scope.xlsx", plan_version_id="plan-other-active", db_path=isolated_db
    )
    assert selected == {
        "status": "failed", "warnings": ["WORKBOOK_EXPORT_HORIZON_UNAVAILABLE"]
    }


def test_export_preserves_staggered_member_periods_without_cartesian_capacity(
    isolated_db: Path, tmp_path: Path
) -> None:
    init_db(quiet=True)
    source = _baseline_workbook(
        tmp_path / "staggered.xlsx",
        start_month="2026-09",
        end_month="2026-10",
        members=[
            ["WD100001", "Alex Example", "LTFTE", "active", None, None, "Engineer", 7, "2026-09-01", "2026-09-30"],
            ["WD100002", "Blair Example", "LTFTE", "active", None, None, "Engineer", 7, "2026-10-01", None],
        ],
        allocations=[
            ["WD100001", "RP-PROJ-001", "2026-09", 0.5],
            ["WD100002", "RP-PROJ-001", "2026-10", 0.5],
        ],
        capacity_rows=[
            ["WD100001", "2026-09", 0.0, 0.0, 0.0],
            ["WD100002", "2026-10", 0.0, 0.0, 0.0],
        ],
    )
    imported = _onboard_workbook(source, db_path=isolated_db)
    assert imported["status"] == "completed", imported
    output = tmp_path / "staggered-export.xlsx"
    assert export_current_state_workbook(output, db_path=isolated_db)["status"] == "success"
    parsed = parse_workbook(output)
    assert {(row.member_key, row.month) for row in parsed.capacity_rows} == {
        ("WD100001", "2026-09"), ("WD100002", "2026-10")
    }


def test_export_keeps_manifest_member_without_a_horizon_period(
    isolated_db: Path, tmp_path: Path
) -> None:
    init_db(quiet=True)
    source = _write_workbook(
        tmp_path / "outside-horizon.xlsx",
        setup_rows=[["FY26 Q4 Baseline", "2026-08-15", "2026-09", "2026-12"]],
        member_rows=[["WD100001", "Alex Example", "LTFTE", "active", None, None, "Engineer", 7, "2027-01-01", None]],
        project_rows=[["RP-PROJ-001", "Project Atlas", "active", 2, "2026-09-01", "2026-12-31"]],
        allocation_rows=[], capacity_rows=[],
    )
    imported = _onboard_workbook(source, db_path=isolated_db)
    assert imported["status"] == "completed", imported
    output = tmp_path / "outside-horizon-export.xlsx"
    result = runner.invoke(app, ["onboarding", "export-workbook", "--output", str(output)])
    assert result.exit_code == 0, result.output
    assert len(result.output.strip().splitlines()) == 1
    assert json.loads(result.output)["status"] == "success"
    parsed = parse_workbook(output)
    assert (parsed.setup_rows[0].start_month, parsed.setup_rows[0].end_month) == ("2026-09", "2026-12")
    assert [member.member_key for member in parsed.members] == ["WD100001"]
    assert parsed.allocations == []
    assert parsed.capacity_rows == []


def test_export_succeeds_when_derived_tables_are_absent(
    isolated_db: Path, tmp_path: Path
) -> None:
    init_db(quiet=True)
    assert _onboard_workbook(
        _baseline_workbook(tmp_path / "publication-mismatch.xlsx", include_capacity=True),
        db_path=isolated_db,
    )["status"] == "completed"
    with sqlite3.connect(isolated_db) as database:
        database.execute("DELETE FROM current_state_staffing_members")
        database.execute("DELETE FROM current_state_staffing_assignments")
        database.execute("DELETE FROM contract_coverage_members")
        database.execute("DELETE FROM resource_capacity_publications")
    output = tmp_path / "source-only.xlsx"
    exported = export_current_state_workbook(output, db_path=isolated_db)
    assert exported["status"] == "success"
    assert validate_workbook(parse_workbook(output)).workbook is not None


@pytest.mark.parametrize("replacement", ["{}", '{"workbook_export_metadata": {}}'])
def test_export_rejects_newest_corrupt_source_receipt_without_fallback(
    replacement: str,
    isolated_db: Path, tmp_path: Path
) -> None:
    init_db(quiet=True)
    assert _onboard_workbook(
        _baseline_workbook(tmp_path / "older-valid.xlsx", include_capacity=True),
        db_path=isolated_db, profile_key="corrupt-source-context",
    )["status"] == "completed"
    assert _onboard_workbook(
        _baseline_workbook(tmp_path / "newer-valid.xlsx", include_capacity=True),
        db_path=isolated_db, profile_key="corrupt-source-context",
    )["status"] == "completed"
    with sqlite3.connect(isolated_db) as database:
        database.execute(
            """UPDATE data_onboarding_runs SET source_preview_json=?
               WHERE run_id=(SELECT run_id FROM data_onboarding_runs
                 WHERE source_type='workbook' AND status='completed'
                 ORDER BY completed_at DESC,created_at DESC,rowid DESC LIMIT 1)""",
            [replacement],
        )
    assert export_current_state_workbook(tmp_path / "legacy.xlsx", db_path=isolated_db) == {
        "status": "failed", "warnings": ["WORKBOOK_EXPORT_SOURCE_CONTEXT_INVALID"]
    }


def test_export_direct_import_without_onboarding_context_fails_closed(
    isolated_db: Path, tmp_path: Path
) -> None:
    init_db(quiet=True)
    assert import_workbook(
        _baseline_workbook(tmp_path / "legacy-direct.xlsx", include_capacity=True),
        db_path=isolated_db,
    )["status"] == "completed"
    assert export_current_state_workbook(
        tmp_path / "legacy-direct-export.xlsx", db_path=isolated_db
    ) == {"status": "failed", "warnings": ["WORKBOOK_EXPORT_HORIZON_UNAVAILABLE"]}


def test_export_previews_and_retains_open_hiref_demand(
    isolated_db: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    init_db(quiet=True)
    assert _onboard_workbook(
        _baseline_workbook(
            tmp_path / "source-roundtrip.xlsx", include_capacity=True,
            members=[["WD100001", "Alex Example", "LTFTE", "active", None, None, "Engineer", 7, "2026-09-01", None, "HIREF-NEXT-001"]],
            allocations=[["WD100001", "RP-PROJ-001", "2026-09", 0.5], [None, "RP-PROJ-001", "2026-09", 0.2, "HIREF-OPEN-001"]],
            extra_sheets={"HIREF Requests": [
                ["HIREF-NEXT-001", "RP-PROJ-001", "new", "2026-09-01", "2026-12-31", "next"],
                ["HIREF-OPEN-001", "RP-PROJ-001", "new", "2026-09-01", "2026-12-31", "open"],
            ]},
        ),
        db_path=isolated_db,
    )["status"] == "completed"
    output = tmp_path / "source-roundtrip-export.xlsx"
    assert export_current_state_workbook(output, db_path=isolated_db)["status"] == "success"
    assert preview_workbook_import(output, db_path=isolated_db)["status"] == "previewed"
    parsed = parse_workbook(output)
    assert [(row.hiref_id, row.project_key, row.allocation) for row in parsed.allocations if row.hiref_id] == [
        ("HIREF-OPEN-001", "RP-PROJ-001", 0.2)
    ]
    roundtrip_db = tmp_path / "roundtrip.db"
    monkeypatch.setattr(settings, "database_path", str(roundtrip_db))
    init_db(quiet=True)
    assert import_workbook(output, db_path=roundtrip_db)["status"] == "completed"
    with sqlite3.connect(roundtrip_db) as database:
        assert database.execute(
            "SELECT COUNT(*) FROM placeholder_monthly_allocations"
        ).fetchone()[0] == 1


def test_hiref_bridge_export_snapshot_fails_closed_when_request_is_corrupted(
    isolated_db: Path, tmp_path: Path
) -> None:
    init_db(quiet=True)
    source = _baseline_workbook(
        tmp_path / "hiref-corruption.xlsx", include_capacity=True,
        members=[["WD100001", "Alex Example", "LTFTE", "active", None, None, "Engineer", 7, "2026-09-01", None, "HIREF-NEXT-001"]],
        allocations=[["WD100001", "RP-PROJ-001", "2026-09", 0.5], [None, "RP-PROJ-001", "2026-09", 0.2, "HIREF-OPEN-001"]],
        extra_sheets={"HIREF Requests": [
            ["HIREF-NEXT-001", "RP-PROJ-001", "new", "2026-09-01", "2026-12-31", "next"],
            ["HIREF-OPEN-001", "RP-PROJ-001", "new", "2026-09-01", "2026-12-31", "open"],
        ]},
    )
    assert import_workbook(source, db_path=isolated_db)["status"] == "completed"
    workforce = source_export_snapshot(db_path=isolated_db)
    bridge = hiref_export_snapshot(
        plan_version_id=workforce["plan_version"]["plan_version_id"],
        member_ids=[row["id"] for row in workforce["members"]], db_path=isolated_db,
    )
    assert bridge["open_demand_allocations"][0][-1] == "HIREF-OPEN-001"
    with sqlite3.connect(isolated_db) as database:
        database.execute("DELETE FROM hiref WHERE id='HIREF-OPEN-001'")
    with pytest.raises(ValueError, match="WORKBOOK_EXPORT_HIREF_REFERENCE_INCOMPLETE"):
        hiref_export_snapshot(
            plan_version_id=workforce["plan_version"]["plan_version_id"],
            member_ids=[row["id"] for row in workforce["members"]], db_path=isolated_db,
        )


def test_hiref_bridge_export_snapshot_accepts_legacy_project_id_storage(
    isolated_db: Path, tmp_path: Path
) -> None:
    init_db(quiet=True)
    source = _baseline_workbook(
        tmp_path / "hiref-legacy-project-id.xlsx", include_capacity=True,
        members=[["WD100001", "Alex Example", "LTFTE", "active", None, None, "Engineer", 7, "2026-09-01", None, "HIREF-NEXT-001"]],
        extra_sheets={"HIREF Requests": [
            ["HIREF-NEXT-001", "RP-PROJ-001", "new", "2026-09-01", "2026-12-31", "next"],
        ]},
    )
    assert import_workbook(source, db_path=isolated_db)["status"] == "completed"
    with sqlite3.connect(isolated_db) as database:
        database.execute("UPDATE hiref SET project='RP-PROJ-001' WHERE id='HIREF-NEXT-001'")
    workforce = source_export_snapshot(db_path=isolated_db)
    snapshot = hiref_export_snapshot(
        plan_version_id=workforce["plan_version"]["plan_version_id"],
        member_ids=[row["id"] for row in workforce["members"]], db_path=isolated_db,
    )
    assert snapshot["hiref_requests"][0][1] == "RP-PROJ-001"


def test_preview_workbook_rejects_unknown_mapping_preset_id(
    isolated_db: Path, tmp_path: Path
) -> None:
    init_db(quiet=True)
    workbook_path = _baseline_workbook(tmp_path / "unknown-preset.xlsx")

    result = preview_workbook_import(
        workbook_path,
        mapping_preset_id="workbook-preset-does-not-exist",
        db_path=isolated_db,
    )

    assert result["status"] == "rejected"
    assert result["source_contract"]["mapping_preset"] == {
        "source_type": "workbook",
        "mapping_preset_id": "workbook-preset-does-not-exist",
        "display_name": None,
        "status": "unknown",
    }
    assert result["source_contract"]["resolution"] is None
    assert result["blockers"] == [
        {
            "severity": "blocker",
            "code": "WORKBOOK_MAPPING_PRESET_UNKNOWN",
            "message": "Unknown workbook mapping preset: workbook-preset-does-not-exist.",
            "location": "mapping_preset_id",
        }
    ]


def test_preview_workbook_rejects_unknown_mapping_preset_before_loading_file(
    isolated_db: Path, tmp_path: Path
) -> None:
    init_db(quiet=True)

    result = preview_workbook_import(
        tmp_path / "missing.xlsx",
        mapping_preset_id="workbook-preset-does-not-exist",
        db_path=isolated_db,
    )

    assert result["status"] == "rejected"
    assert result["blockers"][0]["code"] == "WORKBOOK_MAPPING_PRESET_UNKNOWN"


def test_preview_workbook_tolerates_stale_default_profile_preset(
    isolated_db: Path, tmp_path: Path
) -> None:
    init_db(quiet=True)
    workbook_path = _baseline_workbook(tmp_path / "stale-default.xlsx")

    with sqlite3.connect(isolated_db) as connection:
        connection.execute(
            """
            UPDATE onboarding_profiles
            SET mapping_preset_id=?
            WHERE profile_key='default'
            """,
            ["stale-workbook-preset"],
        )
        connection.commit()

    result = preview_workbook_import(workbook_path, db_path=isolated_db)

    assert result["status"] == "previewed"


def test_parse_workbook_rejects_invalid_sheet_contract(tmp_path: Path) -> None:
    workbook_path = _write_workbook(
        tmp_path / "bad-sheets.xlsx",
        setup_rows=[["FY26 Q4 Baseline", "2026-08-15", "2026-09", "2026-09"]],
        member_rows=[],
        project_rows=[],
        allocation_rows=[],
        capacity_rows=[],
        sheet_names=("Setup", "Members", "Projects", "Allocations", "CapacityX"),
    )
    with pytest.raises(WorkbookParseError, match="WORKBOOK_SHEETS_INVALID"):
        parse_workbook(workbook_path)


def test_validate_workbook_blocks_stfte_without_hiref_fields(tmp_path: Path) -> None:
    workbook_path = _write_workbook(
        tmp_path / "invalid-member.xlsx",
        setup_rows=[["FY26 Q4 Baseline", "2026-08-15", "2026-09", "2026-09"]],
        member_rows=[
            [
                "WD100001",
                "Alex Example",
                "STFTE",
                "active",
                None,
                None,
                "Engineer",
                7,
                "2026-09-01",
                None,
            ]
        ],
        project_rows=[["RP-PROJ-001", "Project Atlas", "active", 2, "2026-09-01", "2026-12-31"]],
        allocation_rows=[["WD100001", "RP-PROJ-001", "2026-09", 0.5]],
        capacity_rows=[],
    )
    result = validate_workbook(parse_workbook(workbook_path))
    assert result.workbook is None
    assert "WORKBOOK_MEMBER_STFTE_HIREF_REQUIRED" in [
        item.code for item in result.blockers
    ]


def test_parse_workbook_rejects_extra_columns(tmp_path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Setup"
    worksheet.append(EXPECTED_HEADERS["Setup"] + ["extra_column"])
    worksheet.append(["FY26 Q4 Baseline", "2026-08-15", "2026-09", "2026-09", "ignored"])
    for name in ("Members", "Projects", "Allocations", "Capacity"):
        tab = workbook.create_sheet(name)
        tab.append(EXPECTED_HEADERS[name])
    workbook_path = tmp_path / "extra-columns.xlsx"
    workbook.save(workbook_path)
    with pytest.raises(WorkbookParseError, match="WORKBOOK_HEADERS_INVALID:Setup"):
        parse_workbook(workbook_path)


def test_validate_workbook_blocks_partial_month_member_effective_dates(
    tmp_path: Path,
) -> None:
    workbook_path = _write_workbook(
        tmp_path / "partial-month.xlsx",
        setup_rows=[["FY26 Q4 Baseline", "2026-08-15", "2026-09", "2026-09"]],
        member_rows=[
            [
                "WD100001",
                "Alex Example",
                "LTFTE",
                "active",
                None,
                None,
                "Engineer",
                7,
                "2026-09-15",
                None,
            ]
        ],
        project_rows=[["RP-PROJ-001", "Project Atlas", "active", 2, "2026-09-01", "2026-12-31"]],
        allocation_rows=[["WD100001", "RP-PROJ-001", "2026-09", 0.5]],
        capacity_rows=[],
    )
    result = validate_workbook(parse_workbook(workbook_path))
    assert "WORKBOOK_MEMBER_EFFECTIVE_START_NOT_MONTH_BOUNDARY" in [
        item.code for item in result.blockers
    ]


def test_preview_workbook_expands_sparse_allocations_and_skips_empty_capacity(
    isolated_db: Path, tmp_path: Path
) -> None:
    init_db(quiet=True)
    workbook_path = _write_workbook(
        tmp_path / "sparse.xlsx",
        setup_rows=[["FY26 Q4 Baseline", "2026-08-15", "2026-09", "2026-10"]],
        member_rows=[
            ["WD100001", "Alex Example", "LTFTE", "active", None, None, "Engineer", 7, "2026-09-01", None],
            ["WD100002", "Blair Example", "LTFTE", "active", None, None, "Engineer", 7, "2026-09-01", None],
        ],
        project_rows=[
            ["RP-PROJ-001", "Project Atlas", "active", 2, "2026-09-01", "2026-12-31"],
            ["RP-PROJ-002", "Project Beacon", "planning", 3, None, "2026-12-31"],
        ],
        allocation_rows=[["WD100001", "RP-PROJ-001", "2026-09", 0.5]],
        capacity_rows=[],
    )
    preview = preview_workbook_import(workbook_path, db_path=isolated_db)
    assert preview["status"] == "previewed"
    assert preview["capacity_package"] is None
    assert preview["counts"]["expanded_allocation_records"] == 8
    assert preview["counts"]["current_state_assignment_records"] == 1
    allocation_map = {
        (
            row["member_id"],
            row["project_id"],
            row["year"],
            row["month"],
        ): row["allocation"]
        for row in preview["workforce_package"]["monthly_allocations"]
    }
    assert allocation_map[("WD100001", "RP-PROJ-001", 2026, 9)] == 0.5
    assert allocation_map[("WD100002", "RP-PROJ-002", 2026, 10)] == 0.0


def test_import_workbook_preserves_explicit_zero_and_missing_capacity_unknown(
    isolated_db: Path, tmp_path: Path
) -> None:
    init_db(quiet=True)
    workbook_path = _write_workbook(
        tmp_path / "capacity.xlsx",
        setup_rows=[["FY26 Q4 Baseline", "2026-08-15", "2026-09", "2026-09"]],
        member_rows=[
            ["WD100001", "Alex Example", "LTFTE", "active", None, None, "Engineer", 7, "2026-09-01", None],
            ["WD100002", "Blair Example", "LTFTE", "active", None, None, "Engineer", 7, "2026-09-01", None],
        ],
        project_rows=[["RP-PROJ-001", "Project Atlas", "active", 2, "2026-09-01", "2026-12-31"]],
        allocation_rows=[
            ["WD100001", "RP-PROJ-001", "2026-09", 0.5],
            ["WD100002", "RP-PROJ-001", "2026-09", 0.2],
        ],
        capacity_rows=[["WD100001", "2026-09", 0.0, 0.0, 0.0]],
    )
    result = import_workbook(workbook_path, db_path=isolated_db)
    plan_version_id = result["plan_version"]["plan_version_id"]
    assert result["status"] == "completed"
    assert result["capacity_result"]["report"]["coverage"]["explicit_zero_count"] == 3

    known = get_effective_capacity(
        "WD100001", 2026, 9, plan_version_id, db_path=isolated_db
    )
    missing = get_effective_capacity(
        "WD100002", 2026, 9, plan_version_id, db_path=isolated_db
    )
    assert known["state"] == "known"
    assert known["effective_capacity"] == 1.0
    assert known["available_capacity"] == 0.5
    assert missing["state"] == "unknown"


def test_import_workbook_supports_hiref_bridge_sections_in_current_onboarding(
    isolated_db: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    init_db(quiet=True)
    workbook_path = _baseline_workbook(
        tmp_path / "hiref-bridge.xlsx",
        members=[
            ["WD100001", "Alex Example", "STFTE", "active", "HIREF-ATLAS-001", "2026-10-15", "Engineer", 7, "2026-09-01", None, None],
            ["WD100002", "Blair Example", "STFTE", "active", "HIREF-CEDAR-001", "2026-10-31", "Engineer", 7, "2026-09-01", None, "HIREF-CEDAR-NEXT"],
        ],
        projects=[
            ["RP-PROJ-001", "Project Atlas", "active", 2, "2026-09-01", "2026-12-31"],
            ["RP-PROJ-002", "Project Beacon", "active", 3, "2026-09-01", "2026-12-31"],
            ["RP-PROJ-003", "Project Cedar", "active", 2, "2026-09-01", "2026-12-31"],
        ],
        allocations=[
            ["WD100001", "RP-PROJ-002", "2026-09", 0.5, None],
            ["WD100002", "RP-PROJ-003", "2026-09", 0.8, None],
            [None, "RP-PROJ-002", "2026-10", 0.5, "HIREF-OPEN-001"],
        ],
        capacity_rows=[
            ["WD100001", "2026-09", 0.0, 0.0, 0.0],
            ["WD100001", "2026-10", 0.0, 0.0, 0.0],
            ["WD100002", "2026-09", 0.0, 0.0, 0.0],
            ["WD100002", "2026-10", 0.0, 0.0, 0.0],
        ],
        start_month="2026-09",
        end_month="2026-10",
        extra_sheets={
            "HIREF Requests": [
                ["HIREF-ATLAS-001", "RP-PROJ-001", "extend", "2026-01-01", "2026-10-15", ""],
                ["HIREF-CEDAR-001", "RP-PROJ-003", "extend", "2026-03-01", "2026-10-31", ""],
                ["HIREF-CEDAR-NEXT", "RP-PROJ-003", "extend", "2026-11-01", "2027-06-30", "Renewal already identified"],
                ["HIREF-OPEN-001", "RP-PROJ-002", "new", "2026-10-01", "2027-03-31", "Open contractor demand"],
            ],
        },
    )

    result = _onboard_workbook(workbook_path, db_path=isolated_db)

    assert result["status"] == "completed"

    with sqlite3.connect(isolated_db) as connection:
        next_hiref = connection.execute(
            "SELECT next_hiref FROM employees WHERE id='WD100002'"
        ).fetchone()
        slots = connection.execute(
            "SELECT id, project FROM hiref ORDER BY id"
        ).fetchall()
        placeholder = connection.execute(
            """
            SELECT placeholder_id, source_system, hiref_id, resource_type, status
            FROM staffing_placeholders
            """
        ).fetchone()
        placeholder_allocation = connection.execute(
            """
            SELECT placeholder_id, project_id, year, month, allocation
            FROM placeholder_monthly_allocations
            """
        ).fetchone()
    assert next_hiref == ("HIREF-CEDAR-NEXT",)
    assert slots == [
        ("HIREF-ATLAS-001", "Project Atlas (RP-PROJ-001)"),
        ("HIREF-CEDAR-001", "Project Cedar (RP-PROJ-003)"),
        ("HIREF-CEDAR-NEXT", "Project Cedar (RP-PROJ-003)"),
        ("HIREF-OPEN-001", "Project Beacon (RP-PROJ-002)"),
    ]
    assert placeholder == (
        "placeholder-hiref-open-001",
        "workbook_onboarding",
        "HIREF-OPEN-001",
        "STFTE",
        "planned",
    )
    assert placeholder_allocation == (
        "placeholder-hiref-open-001",
        "RP-PROJ-002",
        2026,
        10,
        0.5,
    )

    exported = tmp_path / "hiref-bridge-export.xlsx"
    assert export_current_state_workbook(exported, db_path=isolated_db)["status"] == "success"
    parsed_export = parse_workbook(exported)
    assert any(
        row.member_key is None and row.hiref_id == "HIREF-OPEN-001"
        for row in parsed_export.allocations
    )
    second_db = tmp_path / "hiref-roundtrip.db"
    monkeypatch.setattr(settings, "database_path", str(second_db))
    init_db(quiet=True)
    roundtrip = _onboard_workbook(exported, db_path=second_db)
    assert roundtrip["status"] == "completed"
    with sqlite3.connect(second_db) as connection:
        assert connection.execute(
            "SELECT next_hiref FROM employees WHERE id='WD100002'"
        ).fetchone() == ("HIREF-CEDAR-NEXT",)
        assert connection.execute(
            "SELECT hiref_id FROM staffing_placeholders"
        ).fetchone() == ("HIREF-OPEN-001",)
        assert connection.execute(
            "SELECT allocation FROM placeholder_monthly_allocations"
        ).fetchone() == (0.5,)
        assert connection.execute("SELECT COUNT(*) FROM hiref").fetchone() == (4,)

    summary = hiref_management_service.summary(days=90)
    placeholders = hiref_management_service.placeholders()
    review = hiref_management_service.review(days=90)
    assert summary.success is True
    assert placeholders.success is True
    assert review.success is True
    assert summary.data["summary"]["open_placeholders"] == 1
    assert placeholders.data["rows"][0]["slot_registered"] is True
    assert any(
        row["name"] == "Blair Example" and row["next_hiref"] == "HIREF-CEDAR-NEXT"
        for row in review.data["rows"]
    )


def test_preview_workbook_rejects_partial_hiref_snapshot_sections(
    isolated_db: Path, tmp_path: Path
) -> None:
    init_db(quiet=True)
    workbook_path = _baseline_workbook(
        tmp_path / "partial-hiref.xlsx",
        members=[
            ["WD100001", "Alex Example", "STFTE", "active", "HIREF-ATLAS-001", "2026-10-15", "Engineer", 7, "2026-09-01", None, "HIREF-ATLAS-NEXT"]
        ],
    )

    result = preview_workbook_import(workbook_path, db_path=isolated_db)

    assert result["status"] == "rejected"
    assert {
        blocker["code"] for blocker in result["blockers"]
    } >= {
        "WORKBOOK_HIREF_REQUESTS_REQUIRED",
        "WORKBOOK_HIREF_NEXT_REQUEST_NOT_FOUND",
        "WORKBOOK_HIREF_CURRENT_REQUEST_NOT_FOUND",
    }


def test_preview_workbook_rejects_open_hiref_demand_project_mismatch(
    isolated_db: Path, tmp_path: Path
) -> None:
    init_db(quiet=True)
    workbook_path = _baseline_workbook(
        tmp_path / "hiref-demand-mismatch.xlsx",
        projects=[
            ["RP-PROJ-001", "Project Atlas", "active", 2, "2026-09-01", "2026-12-31"],
            ["RP-PROJ-002", "Project Beacon", "active", 3, "2026-09-01", "2026-12-31"],
        ],
        allocations=[
            ["WD100001", "RP-PROJ-001", "2026-09", 0.5, None],
            [None, "RP-PROJ-002", "2026-09", 0.3, "HIREF-ATLAS-001"],
        ],
        members=[
            ["WD100001", "Alex Example", "LTFTE", "active", None, None, "Engineer", 7, "2026-09-01", None, None]
        ],
        extra_sheets={
            "HIREF Requests": [
                ["HIREF-ATLAS-001", "RP-PROJ-001", "new", "2026-09-01", "2026-12-31", ""]
            ]
        },
    )

    result = preview_workbook_import(workbook_path, db_path=isolated_db)

    assert result["status"] == "rejected"
    assert {
        blocker["code"] for blocker in result["blockers"]
    } >= {"WORKBOOK_HIREF_OPEN_DEMAND_PROJECT_MISMATCH"}


def test_import_workbook_empty_hiref_snapshot_clears_stale_bridge_rows(
    isolated_db: Path, tmp_path: Path
) -> None:
    init_db(quiet=True)
    with sqlite3.connect(isolated_db) as connection:
        connection.execute(
            """
            INSERT INTO employees
                (id, wd_id, name, level, status, resource_type, next_hiref)
            VALUES ('WD100001', 'WD100001', 'Alex Example', '7', 'active', 'LTFTE', 'STALE-HIREF')
            """
        )
        connection.execute(
            """
            INSERT INTO projects (id, name, status, priority)
            VALUES ('RP-PROJ-001', 'Project Atlas', 'active', 2)
            """
        )
        connection.execute(
            """
            INSERT INTO plan_versions
                (plan_version_id, version_name, scenario_type, version_status)
            VALUES ('legacy-plan', 'Legacy Plan', 'baseline', 'active')
            """
        )
        connection.execute(
            """
            INSERT INTO hiref (id, project, request_type, start_date, end_date, notes)
            VALUES ('STALE-HIREF', 'Project Atlas (RP-PROJ-001)', 'extend', '2026-01-01', '2026-12-31', '')
            """
        )
        connection.execute(
            """
            INSERT INTO staffing_placeholders
                (placeholder_id, display_name, source_system, hiref_id, status, notes)
            VALUES ('placeholder-stale', 'Stale Placeholder', 'resource_portal', 'STALE-HIREF', 'planned', '')
            """
        )
        connection.execute(
            """
            INSERT INTO placeholder_monthly_allocations
                (placeholder_id, project_id, year, month, allocation, plan_version_id)
            VALUES ('placeholder-stale', 'RP-PROJ-001', 2026, 9, 0.5, 'legacy-plan')
            """
        )
        connection.commit()

    workbook_path = _baseline_workbook(
        tmp_path / "empty-hiref-snapshot.xlsx",
        extra_sheets={
            "HIREF Requests": [],
        },
    )

    result = import_workbook(workbook_path, db_path=isolated_db)

    assert result["status"] == "completed"
    assert result["hiref_bridge_result"]["report"] == {
        "member_next_hiref_rows": 1,
        "hiref_slot_rows": 0,
        "placeholder_rows": 0,
        "placeholder_allocation_rows": 0,
    }
    with sqlite3.connect(isolated_db) as connection:
        employee = connection.execute(
            "SELECT next_hiref FROM employees WHERE id='WD100001'"
        ).fetchone()
        hiref_count = connection.execute("SELECT COUNT(*) FROM hiref").fetchone()
        placeholder_count = connection.execute(
            "SELECT COUNT(*) FROM staffing_placeholders"
        ).fetchone()
        placeholder_allocation_count = connection.execute(
            "SELECT COUNT(*) FROM placeholder_monthly_allocations"
        ).fetchone()
    assert employee == ("",)
    assert hiref_count == (0,)
    assert placeholder_count == (0,)
    assert placeholder_allocation_count == (0,)


def test_import_workbook_rejects_when_workforce_preview_rejects(
    isolated_db: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    init_db(quiet=True)
    monkeypatch.setattr(
        workbook_service,
        "_build_candidate",
        lambda *args, **kwargs: {
            "status": "previewed",
            "workbook_path": "test.xlsx",
            "blockers": [],
            "warnings": [],
            "conflicts": [],
            "plan_version": {
                "plan_version_id": "plan-workbook-test",
                "version_name": "FY26 Q4 Baseline",
            },
            "revision": 1,
            "counts": {},
            "workforce_package": {},
            "current_state_staffing_package": {},
            "capacity_package": None,
        },
    )
    monkeypatch.setattr(
        workbook_service,
        "preview_workforce_import",
        lambda *args, **kwargs: {
            "status": "rejected",
            "failure_code": "WORKFORCE_PLANNING_PACKAGE_REPLAY_CONFLICT",
        },
    )
    result = import_workbook("test.xlsx", db_path=isolated_db)
    assert result["status"] == "rejected"
    assert result["workforce_result"]["status"] == "rejected"
    assert "current_state_staffing_result" not in result
    assert result["capacity_result"] is None


def test_import_workbook_rejection_restores_revision_and_releases_reservation(
    isolated_db: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    init_db(quiet=True)
    workbook_path = _baseline_workbook(tmp_path / "rejected-import.xlsx")
    monkeypatch.setattr(
        workbook_service,
        "preview_workforce_import",
        lambda *args, **kwargs: {
            "status": "rejected",
            "failure_code": "WORKFORCE_PLANNING_PACKAGE_REPLAY_CONFLICT",
        },
    )

    result = import_workbook(workbook_path, db_path=isolated_db)
    assert result["status"] == "rejected"
    assert result["revision"] == 1

    with sqlite3.connect(isolated_db) as connection:
        revision = connection.execute(
            "SELECT current_revision FROM onboarding_profiles WHERE profile_key='default'"
        ).fetchone()
        reservations = connection.execute(
            """
            SELECT status
            FROM data_onboarding_plan_reservations
            WHERE requested_name='FY26 Q4 Baseline'
            """
        ).fetchall()
    assert revision == (0,)
    assert reservations == [("released",)]


def test_confirm_workbook_candidate_reuses_retryable_workforce_session(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    candidate = {
        "status": "previewed",
        "workbook_path": "test.xlsx",
        "blockers": [],
        "warnings": [],
        "conflicts": [],
        "plan_version": {
            "plan_version_id": "plan-workbook-test",
            "version_name": "FY26 Q4 Baseline",
        },
        "revision": 1,
        "counts": {},
        "workforce_package": {},
        "current_state_staffing_package": {
            "dataset_marker": "WORKBOOK_ONBOARDING_V1",
            "package_id": "package-workbook-current-state-test-r1",
            "schema_version": "current-state-staffing-v1",
            "generated_at": "2026-08-15T00:00:00+00:00",
            "source_id": "source-workbook-current-state-staffing",
            "publication_scope": {
                "scope_key": "workbook-current-state-staffing",
                "as_of_date": "2026-08-15",
                "effective_year": 2026,
                "effective_month": 9,
            },
            "manifest": {"member_ids": ["WD100001"], "project_ids": ["RP-PROJ-001"], "assignment_keys": []},
            "members": [
                {
                    "member_id": "WD100001",
                    "display_name": "Alex Example",
                    "status": "active",
                    "role": "Engineer",
                    "level": "7",
                    "resource_type": "LTFTE",
                    "current_hiref_id": None,
                    "hiref_end_date": None,
                }
            ],
            "projects": [
                {
                    "project_id": "RP-PROJ-001",
                    "display_name": "Project Atlas",
                    "status": "active",
                    "priority": 2,
                }
            ],
            "assignments": [],
        },
        "capacity_package": None,
    }
    monkeypatch.setattr(
        workbook_service,
        "preview_workforce_import",
        lambda *args, **kwargs: {
            "status": "retryable",
            "session_id": "workforce-session-1",
        },
    )
    monkeypatch.setattr(
        workbook_service,
        "confirm_workforce_import",
        lambda session_id, **kwargs: {
            "status": "completed",
            "session_id": session_id,
            "idempotent": False,
            "report": {"publication_id": "workforce-publication-1"},
        },
    )
    monkeypatch.setattr(
        workbook_service,
        "preview_current_state_staffing_import",
        lambda *args, **kwargs: {
            "status": "already_completed",
            "session_id": "current-state-session-1",
            "report": {"publication_id": "current-state-publication-1"},
        },
    )
    result = workbook_service.confirm_workbook_candidate(candidate)
    assert result["status"] == "completed"
    assert result["workforce_result"]["session_id"] == "workforce-session-1"
    assert result["workforce_result"]["report"]["publication_id"] == (
        "workforce-publication-1"
    )
    assert result["current_state_staffing_result"]["status"] == "completed"


def test_confirm_workbook_candidate_accepts_already_completed_workforce_session(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    candidate = {
        "status": "previewed",
        "workbook_path": "test.xlsx",
        "blockers": [],
        "warnings": [],
        "conflicts": [],
        "plan_version": {
            "plan_version_id": "plan-workbook-test",
            "version_name": "FY26 Q4 Baseline",
        },
        "revision": 1,
        "counts": {},
        "workforce_package": {},
        "current_state_staffing_package": {
            "dataset_marker": "WORKBOOK_ONBOARDING_V1",
            "package_id": "package-workbook-current-state-test-r1",
            "schema_version": "current-state-staffing-v1",
            "generated_at": "2026-08-15T00:00:00+00:00",
            "source_id": "source-workbook-current-state-staffing",
            "publication_scope": {
                "scope_key": "workbook-current-state-staffing",
                "as_of_date": "2026-08-15",
                "effective_year": 2026,
                "effective_month": 9,
            },
            "manifest": {"member_ids": ["WD100001"], "project_ids": ["RP-PROJ-001"], "assignment_keys": []},
            "members": [
                {
                    "member_id": "WD100001",
                    "display_name": "Alex Example",
                    "status": "active",
                    "role": "Engineer",
                    "level": "7",
                    "resource_type": "LTFTE",
                    "current_hiref_id": None,
                    "hiref_end_date": None,
                }
            ],
            "projects": [
                {
                    "project_id": "RP-PROJ-001",
                    "display_name": "Project Atlas",
                    "status": "active",
                    "priority": 2,
                }
            ],
            "assignments": [],
        },
        "capacity_package": None,
    }
    monkeypatch.setattr(
        workbook_service,
        "preview_workforce_import",
        lambda *args, **kwargs: {
            "status": "already_completed",
            "session_id": "workforce-session-1",
            "report": {"publication_id": "workforce-publication-1"},
        },
    )
    monkeypatch.setattr(
        workbook_service,
        "preview_current_state_staffing_import",
        lambda *args, **kwargs: {
            "status": "already_completed",
            "session_id": "current-state-session-1",
            "report": {"publication_id": "current-state-publication-1"},
        },
    )
    result = workbook_service.confirm_workbook_candidate(candidate)
    assert result["status"] == "completed"
    assert result["workforce_result"] == {
        "status": "completed",
        "session_id": "workforce-session-1",
        "idempotent": True,
        "report": {"publication_id": "workforce-publication-1"},
    }
    assert result["current_state_staffing_result"]["status"] == "completed"


def test_advance_profile_revision_serializes_concurrent_claims(isolated_db: Path) -> None:
    init_db(quiet=True)
    barrier = threading.Barrier(2)
    results: list[tuple[str, int]] = []
    errors: list[BaseException] = []

    def worker() -> None:
        try:
            barrier.wait()
            results.append(
                (
                    "ok",
                    advance_profile_revision(profile_key="default", db_path=isolated_db),
                )
            )
        except BaseException as exc:  # pragma: no cover - failure path assertion below
            errors.append(exc)

    threads = [threading.Thread(target=worker) for _ in range(2)]
    for thread in threads:
        thread.start()
    for thread in threads:
        thread.join()

    assert errors == []
    assert sorted(revision for _, revision in results) == [1, 2]


def test_resolve_plan_identity_serializes_concurrent_reservations(
    isolated_db: Path,
) -> None:
    init_db(quiet=True)
    barrier = threading.Barrier(2)
    results: list[dict[str, str]] = []
    errors: list[BaseException] = []

    def worker(profile_key: str, run_id: str) -> None:
        try:
            barrier.wait()
            results.append(
                resolve_plan_identity(
                    "FY26 Q4 Baseline",
                    profile_key=profile_key,
                    run_id=run_id,
                    db_path=isolated_db,
                )
            )
        except BaseException as exc:  # pragma: no cover - failure path assertion below
            errors.append(exc)

    threads = [
        threading.Thread(target=worker, args=("profile-a", "run-a")),
        threading.Thread(target=worker, args=("profile-b", "run-b")),
    ]
    for thread in threads:
        thread.start()
    for thread in threads:
        thread.join()

    assert errors == []
    assert {item["version_name"] for item in results} == {
        "FY26 Q4 Baseline",
        "FY26 Q4 Baseline-v2",
    }
    assert {item["plan_version_id"] for item in results} == {
        "plan-workbook-fy26-q4-baseline",
        "plan-workbook-fy26-q4-baseline-v2",
    }
    with sqlite3.connect(isolated_db) as connection:
        reservations = connection.execute(
            """
            SELECT run_id, version_name, plan_version_id, status
            FROM data_onboarding_plan_reservations
            ORDER BY run_id
            """
        ).fetchall()
    assert len(reservations) == 2
    assert set(reservations).issubset(
        {
        ("run-a", "FY26 Q4 Baseline", "plan-workbook-fy26-q4-baseline", "reserved"),
        ("run-a", "FY26 Q4 Baseline-v2", "plan-workbook-fy26-q4-baseline-v2", "reserved"),
        ("run-b", "FY26 Q4 Baseline", "plan-workbook-fy26-q4-baseline", "reserved"),
        ("run-b", "FY26 Q4 Baseline-v2", "plan-workbook-fy26-q4-baseline-v2", "reserved"),
        }
    )


def test_reimport_deduplicates_plan_version_name_and_replaces_current_baseline(
    isolated_db: Path, tmp_path: Path
) -> None:
    init_db(quiet=True)
    first_path = _baseline_workbook(tmp_path / "first.xlsx")
    first = import_workbook(first_path, db_path=isolated_db)
    second_path = _baseline_workbook(tmp_path / "second.xlsx")
    second = import_workbook(second_path, db_path=isolated_db)

    assert first["status"] == "completed"
    assert second["status"] == "completed"
    assert first["plan_version"]["version_name"] == "FY26 Q4 Baseline"
    assert second["plan_version"]["version_name"] == "FY26 Q4 Baseline-v2"

    current = workforce_repository.current_publication(db_path=isolated_db)
    package = json.loads(current["package_json"])
    assert package["plan_versions"][0]["version_name"] == "FY26 Q4 Baseline-v2"

    with sqlite3.connect(isolated_db) as connection:
        statuses = dict(
            connection.execute(
                "SELECT version_name, version_status FROM plan_versions"
            ).fetchall()
        )
    assert statuses == {
        "FY26 Q4 Baseline": "archived",
        "FY26 Q4 Baseline-v2": "active",
    }


def test_workbook_preview_blocks_unabsorbed_confirmed_adjustment_and_allows_absorbed(
    isolated_db: Path, tmp_path: Path
) -> None:
    init_db(quiet=True)
    baseline_path = _baseline_workbook(tmp_path / "baseline.xlsx", allocation_value=0.5)
    baseline = import_workbook(baseline_path, db_path=isolated_db)
    assert baseline["status"] == "completed"
    current_plan_id = baseline["plan_version"]["plan_version_id"]

    _confirm_staffing_adjustment(
        project_id="RP-PROJ-001",
        member_id="WD100001",
        plan_version_id=current_plan_id,
        month="2026-09",
        allocation=0.1,
    )

    conflicting = preview_workbook_import(
        _baseline_workbook(tmp_path / "conflict.xlsx", allocation_value=0.5),
        db_path=isolated_db,
    )
    absorbed = preview_workbook_import(
        _baseline_workbook(tmp_path / "absorbed.xlsx", allocation_value=0.6),
        db_path=isolated_db,
    )

    assert conflicting["status"] == "rejected"
    assert len(conflicting["conflicts"]) == 1
    assert conflicting["conflicts"][0]["member_id"] == "WD100001"
    assert conflicting["conflicts"][0]["project_id"] == "RP-PROJ-001"
    assert conflicting["conflicts"][0]["plan_version_id"] == current_plan_id
    assert conflicting["conflicts"][0]["baseline_allocation"] == 0.5
    assert conflicting["conflicts"][0]["effective_allocation"] == 0.6
    assert conflicting["conflicts"][0]["candidate_allocation"] == 0.5
    assert (
        conflicting["conflicts"][0]["conflict_code"]
        == "WORKBOOK_ONBOARDING_CONFIRMED_ADJUSTMENT_CONFLICT"
    )
    assert absorbed["status"] == "previewed"
    assert absorbed["conflicts"] == []
