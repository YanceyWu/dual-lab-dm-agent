from __future__ import annotations

import csv
import json
import sqlite3
from pathlib import Path

from openpyxl import load_workbook
from typer.testing import CliRunner

from pm_agent.cli.app import app
from pm_agent.database.bootstrap import main as init_db

runner = CliRunner()


def _invoke(*args: str):
    return runner.invoke(app, list(args))


def _payload(result) -> dict:
    return json.loads(result.stdout)


def _preview_exported(path: Path, source_type: str, profile_key: str) -> dict:
    saved = _invoke(
        "onboarding", "profile", "save", "--profile-key", profile_key,
        "--source-type", source_type, "--file", str(path),
    )
    assert saved.exit_code == 0, saved.output
    preview = _invoke("onboarding", "preview", "--profile-key", profile_key)
    assert preview.exit_code == 0, preview.output
    return _payload(preview)


def test_export_source_artifacts_use_existing_handlers_and_exclude_derived_rows(
    isolated_db: Path, tmp_path: Path
) -> None:
    init_db(quiet=True)
    with sqlite3.connect(isolated_db) as connection:
        connection.execute(
            "INSERT INTO projects (id, name, status, priority) VALUES ('project-001', 'Project One', 'active', 1)"
        )
        connection.execute(
            """
            INSERT INTO project_profiles
                (project_id, phase, phase_detail, priority_tier, is_focus,
                 objective, milestones, key_risks, stakeholders, special_rules)
            VALUES ('project-001', 'delivery', 'steady', 1, 1, 'Synthetic objective',
                    '[{\"date\":\"2026-08-10\",\"name\":\"Gate\"}]',
                    '[{\"risk\":\"Synthetic risk\",\"level\":\"high\"}]',
                    '[\"Synthetic stakeholder\"]', 'Synthetic rule')
            """
        )
        connection.execute(
            """
            INSERT INTO jira_board_configs
                (id, name, project_key, base_jql, version_name_pattern, board_id,
                 board_url, pm_project_id, active, issues_use_base_jql, notes)
            VALUES ('board-001', 'Board One', 'ONE', 'project = ONE', '%ONE%', '123',
                    'https://local.example/board/123', 'project-001', 1, 0, 'Synthetic note')
            """
        )
        connection.execute(
            "INSERT INTO jira_health_snapshots (board_id, snapshot_date) VALUES ('board-001', '2026-08-09')"
        )
        connection.execute(
            """
            INSERT INTO confluence_pages
                (id, board_id, title, page_type, last_synced, last_modified, content_summary)
            VALUES ('page-001', 'board-001', 'Status', 'status_page', '2026-08-09',
                    '2026-08-08', 'derived connector content')
            """
        )
        connection.execute(
            "INSERT INTO confluence_status_snapshots (board_id, page_id, snapshot_date) VALUES ('board-001', 'page-001', '2026-08-09')"
        )
        connection.commit()

    profile_path = tmp_path / "project_profiles.xlsx"
    jira_path = tmp_path / "jira_boards.csv"
    confluence_path = tmp_path / "confluence_pages.csv"
    for source_type, path in (
        ("project-profile-workbook", profile_path),
        ("jira-board-registry-csv", jira_path),
        ("confluence-page-registry-csv", confluence_path),
    ):
        result = _invoke(
            "onboarding", "export-source", "--source-type", source_type,
            "--output", str(path),
        )
        assert result.exit_code == 0, result.output
        payload = _payload(result)
        assert payload["status"] == "exported"
        assert payload["source_type"] == source_type
        assert payload["output_path"] == str(path.resolve())
        assert len(payload["sha256"]) == 64

    workbook = load_workbook(profile_path, data_only=True)
    assert workbook["Project Profiles"].cell(row=5, column=12).value == "project-001"
    assert _preview_exported(profile_path, "project-profile-workbook", "profiles")["status"] in {
        "previewed", "already_completed"
    }
    assert _preview_exported(jira_path, "jira-board-registry-csv", "jira")["status"] in {
        "previewed", "already_completed"
    }
    confluence_preview = _preview_exported(
        confluence_path, "confluence-page-registry-csv", "confluence"
    )
    assert confluence_preview["status"] in {"previewed", "already_completed"}
    with confluence_path.open(encoding="utf-8", newline="") as handle:
        row = next(csv.DictReader(handle))
    assert row["id"] == "page-001"
    assert row["last_synced"] == ""
    assert row["last_modified"] == ""
    assert row["content_summary"] == ""


def test_export_source_protects_existing_paths_and_supports_explicit_overwrite(
    isolated_db: Path, tmp_path: Path
) -> None:
    init_db(quiet=True)
    output = tmp_path / "jira_boards.csv"
    output.write_text("keep", encoding="utf-8")

    refused = _invoke(
        "onboarding", "export-source", "--source-type", "jira-board-registry-csv",
        "--output", str(output),
    )
    assert refused.exit_code == 2
    assert _payload(refused)["warnings"] == ["DATA_ONBOARDING_EXPORT_OUTPUT_EXISTS"]
    assert output.read_text(encoding="utf-8") == "keep"

    overwritten = _invoke(
        "onboarding", "export-source", "--source-type", "jira-board-registry-csv",
        "--output", str(output), "--overwrite",
    )
    assert overwritten.exit_code == 0, overwritten.output
    assert _payload(overwritten)["status"] == "exported"
    with output.open(encoding="utf-8", newline="") as handle:
        assert next(csv.reader(handle)) == [
            "id", "name", "project_key", "base_jql", "version_name_pattern",
            "board_id", "board_url", "pm_project_id", "active",
            "issues_use_base_jql", "notes",
        ]


def test_export_source_rejects_unknown_type_with_single_json_result(tmp_path: Path) -> None:
    result = _invoke(
        "onboarding", "export-source", "--source-type", "unsupported",
        "--output", str(tmp_path / "ignored.csv"),
    )
    assert result.exit_code == 2
    assert _payload(result) == {
        "status": "failed",
        "source_type": "unsupported",
        "warnings": ["DATA_ONBOARDING_EXPORT_SOURCE_TYPE_UNSUPPORTED"],
    }


def test_export_source_rejects_a_path_incompatible_with_its_existing_contract(
    tmp_path: Path,
) -> None:
    result = _invoke(
        "onboarding", "export-source", "--source-type", "project-profile-workbook",
        "--output", str(tmp_path / "project_profiles.csv"),
    )
    assert result.exit_code == 2
    assert _payload(result)["warnings"] == ["DATA_ONBOARDING_EXPORT_OUTPUT_SUFFIX_INVALID"]


def test_project_profile_export_fails_closed_for_malformed_json_and_noncanonical_flags(
    isolated_db: Path, tmp_path: Path
) -> None:
    init_db(quiet=True)
    with sqlite3.connect(isolated_db) as connection:
        connection.execute(
            "INSERT INTO projects (id, name, status, priority) VALUES ('project-001', 'Project One', 'active', 1)"
        )
        connection.execute(
            """
            INSERT INTO project_profiles (project_id, phase, priority_tier, is_focus, milestones)
            VALUES ('project-001', 'delivery', 1, 1, '{not-json')
            """
        )
        connection.commit()
    output = tmp_path / "profiles.xlsx"
    malformed = _invoke(
        "onboarding", "export-source", "--source-type", "project-profile-workbook",
        "--output", str(output),
    )
    assert malformed.exit_code == 2
    assert _payload(malformed)["warnings"] == ["PROJECT_PROFILE_EXPORT_SOURCE_UNREPRESENTABLE"]
    assert not output.exists()

    with sqlite3.connect(isolated_db) as connection:
        connection.execute("UPDATE project_profiles SET milestones = '[]', is_focus = 2")
        connection.commit()
    noncanonical = _invoke(
        "onboarding", "export-source", "--source-type", "project-profile-workbook",
        "--output", str(output),
    )
    assert noncanonical.exit_code == 2
    assert _payload(noncanonical)["warnings"] == ["PROJECT_PROFILE_EXPORT_SOURCE_UNREPRESENTABLE"]
    assert not output.exists()


def test_jira_registry_export_fails_closed_for_boolean_and_parser_normalization_loss(
    isolated_db: Path, tmp_path: Path
) -> None:
    init_db(quiet=True)
    with sqlite3.connect(isolated_db) as connection:
        connection.execute(
            """
            INSERT INTO jira_board_configs (id, name, project_key, base_jql, active, issues_use_base_jql)
            VALUES ('board-001', 'Board One', 'ONE', 'project = ONE', 2, 0)
            """
        )
        connection.commit()
    output = tmp_path / "jira.csv"
    invalid_boolean = _invoke(
        "onboarding", "export-source", "--source-type", "jira-board-registry-csv",
        "--output", str(output),
    )
    assert invalid_boolean.exit_code == 2
    assert _payload(invalid_boolean)["warnings"] == ["JIRA_BOARD_REGISTRY_EXPORT_SOURCE_UNREPRESENTABLE"]
    assert not output.exists()

    with sqlite3.connect(isolated_db) as connection:
        connection.execute("UPDATE jira_board_configs SET active = 1, project_key = 'one'")
        connection.commit()
    normalized = _invoke(
        "onboarding", "export-source", "--source-type", "jira-board-registry-csv",
        "--output", str(output),
    )
    assert normalized.exit_code == 2
    assert _payload(normalized)["warnings"] == ["JIRA_BOARD_REGISTRY_EXPORT_SOURCE_UNREPRESENTABLE"]
    assert not output.exists()


def test_confluence_registry_export_fails_closed_when_import_would_collapse_or_change_rows(
    isolated_db: Path, tmp_path: Path
) -> None:
    init_db(quiet=True)
    with sqlite3.connect(isolated_db) as connection:
        connection.executemany(
            "INSERT INTO confluence_pages (id, board_id, title, page_type) VALUES (?, 'board-001', 'Status', 'status_page')",
            [("page-001",), ("page-002",)],
        )
        connection.commit()
    output = tmp_path / "confluence.csv"
    collapsed = _invoke(
        "onboarding", "export-source", "--source-type", "confluence-page-registry-csv",
        "--output", str(output),
    )
    assert collapsed.exit_code == 2
    assert _payload(collapsed)["warnings"] == ["CONFLUENCE_PAGE_REGISTRY_EXPORT_SOURCE_UNREPRESENTABLE"]
    assert not output.exists()

    with sqlite3.connect(isolated_db) as connection:
        connection.execute("DELETE FROM confluence_pages WHERE id = 'page-002'")
        connection.execute("UPDATE confluence_pages SET page_type = '' WHERE id = 'page-001'")
        connection.commit()
    changed = _invoke(
        "onboarding", "export-source", "--source-type", "confluence-page-registry-csv",
        "--output", str(output),
    )
    assert changed.exit_code == 2
    assert _payload(changed)["warnings"] == ["CONFLUENCE_PAGE_REGISTRY_EXPORT_SOURCE_UNREPRESENTABLE"]
    assert not output.exists()
