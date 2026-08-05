from __future__ import annotations

"""
scripts/import_from_excel.py — Legacy Distribution Excel import helper.

Usage:
    python3 scripts/import_from_excel.py --file "data-feed/distribution.xlsx"

What it does:
1. Wipes existing employees, projects, assignments (clean slate)
2. Parses every staff member and project from the Excel
3. Stores employees by canonical source ID
4. Routes HIREF-only / placeholder rows into staffing_placeholders
5. Creates assignments using the current month's allocation value

Current-state staffing readers no longer treat those assignment rows as the
authoritative workload source. Use the workbook onboarding entrypoint when you
need canonical current-state staffing publication behavior.
"""

import re
import sys
import sqlite3
import argparse
from pathlib import Path
from datetime import date

sys.path.insert(0, str(Path(__file__).parent.parent))
from pm_agent.database.bootstrap import main as init_db
from pm_agent.rules.identity import (
    build_placeholder_id,
    derive_workday_id,
    is_placeholder_identifier,
    normalize_resource_portal_id,
)
from pm_agent.database import repository
from pm_agent.config import settings

try:
    import openpyxl
except ImportError:
    print("Run: pip3 install openpyxl")
    sys.exit(1)


# ──────────────────────────────────────────────
# Parsing helpers
# ──────────────────────────────────────────────

def parse_staff(raw: str) -> dict[str, str | bool | None]:
    """
    Returns the canonical identity payload for one Resource Portal staff cell.

    Examples use fictional identities:
      'Alex Example - 990001'
      'Blair Example - CQA990002'
      'HIREF-SYN-001'
    """
    raw = raw.strip()
    display_name = raw
    zh_name = ''
    source_identifier = ''

    # Pattern: English (Chinese) - ID
    m = re.match(r'^(.+?)\s*[（(]([^）)]+)[）)]\s*-\s*(.+)$', raw)
    if m:
        en_name = m.group(1).strip()
        zh_name = m.group(2).strip()
        source_identifier = m.group(3).strip()
        display_name = f"{en_name} ({zh_name})"
    else:
        # Pattern: English - ID
        m = re.match(r'^(.+?)\s+-\s+(.+)$', raw)
        if m:
            display_name = m.group(1).strip()
            source_identifier = m.group(2).strip()

    normalized_source_id = normalize_resource_portal_id(source_identifier or raw)
    lower_raw = raw.lower()
    is_placeholder = (
        is_placeholder_identifier(normalized_source_id)
        or 'to be hired' in lower_raw
    )

    if is_placeholder:
        hiref_id = ''
        if source_identifier and is_placeholder_identifier(source_identifier):
            hiref_id = source_identifier.strip()
        elif is_placeholder_identifier(raw):
            hiref_id = raw.strip()
        return {
            'display_name': display_name,
            'zh_name': zh_name,
            'source_employee_id': '',
            'workday_id': None,
            'is_placeholder': True,
            'placeholder_id': build_placeholder_id(source_identifier or raw, display_name),
            'hiref_id': hiref_id,
        }

    workday_id = derive_workday_id(normalized_source_id)
    if not workday_id:
        raise ValueError(f"Cannot derive Workday ID from staff cell: {raw}")

    return {
        'display_name': display_name,
        'zh_name': zh_name,
        'source_employee_id': normalized_source_id,
        'workday_id': workday_id,
        'is_placeholder': False,
        'placeholder_id': None,
        'hiref_id': '',
    }


def parse_project(raw: str) -> tuple[str, str, str]:
    """
    Returns (display_name, id_slug, jira_code)

    Example:
      'Project Atlas (9901001)' →
      ('Project Atlas', 'project-atlas-9901001', '9901001')
    """
    raw = raw.strip()
    m = re.match(r'^(.+?)\s*\((\d+)\)\s*$', raw)
    if m:
        name = m.group(1).strip()
        code = m.group(2).strip()
        slug = re.sub(r'[^a-z0-9]+', '-', name.lower()).strip('-') + f"-{code}"
        return name, slug, code
    slug = re.sub(r'[^a-z0-9]+', '-', raw.lower()).strip('-')
    return raw, slug, ''


# ──────────────────────────────────────────────
# Main import logic
# ──────────────────────────────────────────────

MONTHS = ['Jan', 'Feb', 'Mar', 'Apr', 'May', 'Jun',
          'Jul', 'Aug', 'Sep', 'Oct', 'Nov', 'Dec']

# Current month index (0-based) — May = 4
CURRENT_MONTH_IDX = date.today().month - 1


def load_excel(file_path: str) -> list[dict]:
    """Read the Distribution sheet and return clean records."""
    wb = openpyxl.load_workbook(file_path)
    ws = _find_distribution_sheet(wb)

    records = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        staff_raw   = row[3]
        project_raw = row[4]
        allocs      = row[6:18]  # Jan..Dec columns
        status      = row[22]    # 'Active'|'Inactive'

        if not staff_raw or status != 'Active':
            continue

        records.append({
            'staff_raw':   str(staff_raw),
            'project_raw': str(project_raw) if project_raw else None,
            'allocs':      allocs,  # tuple of 12 values
        })
    return records


def _find_distribution_sheet(wb):
    """Locate the distribution export by its field contract, not a team name."""
    expected_months = MONTHS
    for ws in wb.worksheets:
        staff_header = str(ws.cell(row=1, column=4).value or "").strip().lower()
        project_header = str(ws.cell(row=1, column=5).value or "").strip().lower()
        month_headers = [
            str(ws.cell(row=1, column=column).value or "").strip()
            for column in range(7, 19)
        ]
        status_header = str(ws.cell(row=1, column=23).value or "").strip().lower()
        if (
            staff_header == "staff"
            and project_header == "project"
            and month_headers == expected_months
            and status_header == "status"
        ):
            return ws

    raise ValueError(
        "No distribution worksheet matches the required Staff, Project, "
        "Jan-Dec, and Status column contract."
    )


def import_data(file_path: str, dry_run: bool = False) -> None:
    init_db()
    default_plan_version_id = "baseline-imported-2026"

    records = load_excel(file_path)
    print(f"Loaded {len(records)} active rows from Excel")

    # ── Build unique staff / placeholders / projects ──
    employee_map: dict[str, dict] = {}     # workday_id → member dict
    placeholder_map: dict[str, dict] = {}  # placeholder_id → placeholder dict
    project_map: dict[str, dict] = {}      # slug → project dict

    for rec in records:
        staff = parse_staff(rec['staff_raw'])
        if staff['is_placeholder']:
            placeholder_id = str(staff['placeholder_id'])
            if placeholder_id not in placeholder_map:
                placeholder_map[placeholder_id] = {
                    'placeholder_id': placeholder_id,
                    'display_name': staff['display_name'],
                    'source_employee_id': staff['source_employee_id'] or '',
                    'hiref_id': staff['hiref_id'] or '',
                    'resource_type': '',
                    'status': 'planned',
                    'notes': 'Imported from Excel. Placeholder hire.',
                    'metadata': '{}',
                }
        else:
            workday_id = str(staff['workday_id'])
            if workday_id not in employee_map:
                employee_map[workday_id] = {
                    'id': workday_id,
                    'wd_id': workday_id,
                    'source_employee_id': staff['source_employee_id'],
                    'name': staff['display_name'],
                    'email': None,
                    'role': 'TBC',
                    'level': 'mid',
                    'team': None,
                    'lead_id': None,
                    'max_parallel': 3,
                    'status': 'active',
                    'notes': 'Imported from Excel.',
                    'skills': '{}',
                    'metadata': '{}',
                }

        # Project
        if rec['project_raw']:
            name_p, slug_p, code = parse_project(rec['project_raw'])
            if slug_p not in project_map:
                project_map[slug_p] = {
                    'id':          slug_p,
                    'name':        name_p,
                    'jira_key':    code,
                    'status':      'active',
                    'priority':    3,
                    'lead_id':     None,
                    'tech_stack':  '[]',
                    'start_date':  '2026-01-01',
                    'target_end':  None,
                    'actual_end':  None,
                    'notes':       '',
                }

    print(f"  → {len(employee_map)} unique employees")
    print(f"  → {len(placeholder_map)} unique placeholders")
    print(f"  → {len(project_map)} unique projects")

    if dry_run:
        print("[DRY RUN] Stopping before DB writes.")
        return

    run_id = repository.start_sync_run(
        "import-resource-portal",
        run_type="file-import",
        target_tables=[
            "employees",
            "employee_external_ids",
            "staffing_placeholders",
            "projects",
            "project_profiles",
            "assignments",
            "monthly_allocations",
            "placeholder_monthly_allocations",
            "plan_versions",
        ],
        artifact_path=str(Path(file_path).resolve()),
        triggered_by="cli",
        notes="Distribution Excel import",
    )

    db_path = Path(settings.database_path)
    con = None
    assignment_count = 0
    monthly_count = 0
    placeholder_monthly_count = 0
    ext_count = 0
    override_count = 0
    try:
        con = sqlite3.connect(db_path)
        con.execute("PRAGMA foreign_keys = OFF")  # allow clean wipe

        # ── Wipe existing sample data ──
        con.execute("DELETE FROM project_profiles")
        con.execute(
            """
            DELETE FROM placeholder_monthly_allocations
            WHERE COALESCE(plan_version_id, '') IN ('', ?)
            """,
            [default_plan_version_id],
        )
        con.execute(
            """
            DELETE FROM monthly_allocations
            WHERE COALESCE(plan_version_id, '') IN ('', ?)
            """,
            [default_plan_version_id],
        )
        con.execute("DELETE FROM assignments")
        con.execute("DELETE FROM decision_log")
        con.execute("DELETE FROM action_items")
        con.execute("DELETE FROM employee_external_ids")
        con.execute("DELETE FROM staffing_placeholders")
        con.execute(
            "DELETE FROM plan_versions WHERE plan_version_id = ?",
            [default_plan_version_id],
        )
        con.execute("DELETE FROM employees")
        con.execute("DELETE FROM projects")
        con.commit()
        print("  ✓ Cleared existing data")

        # ── Insert employees ──
        for m in employee_map.values():
            con.execute("""
                INSERT OR REPLACE INTO employees
                (id, wd_id, name, email, role, level, team, lead_id, max_parallel,
                 status, notes, skills, metadata)
                VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)
            """, [m['id'], m['wd_id'], m['name'], m['email'], m['role'], m['level'],
                  m['team'], m['lead_id'], m['max_parallel'],
                  m['status'], m['notes'], m['skills'], m['metadata']])
        con.commit()
        print(f"  ✓ Inserted {len(employee_map)} employees")

        # ── Insert staffing_placeholders (v1.6) ──
        for p in placeholder_map.values():
            con.execute("""
                INSERT OR REPLACE INTO staffing_placeholders
                    (placeholder_id, display_name, source_system, source_employee_id, hiref_id,
                     resource_type, status, notes, metadata)
                VALUES (?, ?, 'resource_portal', ?, ?, ?, ?, ?, ?)
            """, [
                p['placeholder_id'],
                p['display_name'],
                p['source_employee_id'],
                p['hiref_id'],
                p['resource_type'],
                p['status'],
                p['notes'],
                p['metadata'],
            ])
        con.commit()
        print(f"  ✓ Inserted {len(placeholder_map)} staffing placeholders")

        # ── Insert employee_external_ids (v1.6) ──
        for m in employee_map.values():
            emp_id = m["id"]
            name = m["name"]
            source_employee_id = m["source_employee_id"] or emp_id
            con.execute("""
                INSERT OR REPLACE INTO employee_external_ids
                    (employee_id, system_name, id_type, external_id, external_name)
                VALUES (?, 'resource_portal', 'employee_id', ?, ?)
            """, [emp_id, source_employee_id, name])
            ext_count += 1
        con.commit()
        print(f"  ✓ Seeded {ext_count} employee_external_ids rows")

        # ── Insert projects ──
        for p in project_map.values():
            con.execute("""
                INSERT OR REPLACE INTO projects
                (id, name, jira_key, status, priority, lead_id,
                 tech_stack, start_date, target_end, actual_end, notes)
                VALUES (?,?,?,?,?,?,?,?,?,?,?)
            """, [p['id'], p['name'], p['jira_key'], p['status'], p['priority'],
                  p['lead_id'], p['tech_stack'], p['start_date'],
                  p['target_end'], p['actual_end'], p['notes']])
        con.commit()
        print(f"  ✓ Inserted {len(project_map)} projects")

        # ── Insert default plan version (v1.5) ──
        con.execute("""
            INSERT OR REPLACE INTO plan_versions
                (plan_version_id, version_name, scenario_type, as_of_date, version_status, notes)
            VALUES (?, 'Imported Baseline 2026', 'baseline', date('now'), 'active',
                    'Created by import_from_excel.py')
        """, [default_plan_version_id])
        con.execute("""
            DELETE FROM monthly_allocations
            WHERE employee_id NOT IN (SELECT id FROM employees)
               OR project_id NOT IN (SELECT id FROM projects)
        """)
        con.execute("""
            DELETE FROM placeholder_monthly_allocations
            WHERE placeholder_id NOT IN (SELECT placeholder_id FROM staffing_placeholders)
               OR project_id NOT IN (SELECT id FROM projects)
        """)
        con.execute("""
            UPDATE project_snapshots
            SET staffing_scenario_id = NULL
            WHERE staffing_scenario_id IS NOT NULL
              AND staffing_scenario_id NOT IN (SELECT plan_version_id FROM plan_versions)
        """)
        con.execute("""
            DELETE FROM project_snapshots
            WHERE project_id NOT IN (SELECT id FROM projects)
        """)
        con.commit()
        print(f"  ✓ Created plan version: {default_plan_version_id}")

        # ── Insert assignments (current month allocation > 0) ──
        # Also populate monthly_allocations for all 12 months
        seen_assign = set()

        for rec in records:
            staff = parse_staff(rec['staff_raw'])
            if not rec['project_raw']:
                continue

            _, proj_slug, _ = parse_project(rec['project_raw'])
            alloc_values = rec['allocs']

            if staff['is_placeholder']:
                placeholder_id = str(staff['placeholder_id'])
                for m_idx, alloc_val in enumerate(alloc_values):
                    if alloc_val is not None and alloc_val > 0:
                        con.execute("""
                            INSERT OR REPLACE INTO placeholder_monthly_allocations
                                (placeholder_id, project_id, year, month, allocation, plan_version_id)
                            VALUES (?, ?, 2026, ?, ?, ?)
                        """, [
                            placeholder_id,
                            proj_slug,
                            m_idx + 1,
                            float(alloc_val),
                            default_plan_version_id,
                        ])
                        placeholder_monthly_count += 1
                continue

            workday_id = str(staff['workday_id'])

            # Current assignment
            alloc = alloc_values[CURRENT_MONTH_IDX]
            if alloc and alloc > 0:
                key = (workday_id, proj_slug)
                if key not in seen_assign:
                    seen_assign.add(key)
                    con.execute("""
                        INSERT OR IGNORE INTO assignments
                        (employee_id, project_id, role, allocation, start_date, status)
                        VALUES (?, ?, 'member', ?, date('now'), 'active')
                    """, [workday_id, proj_slug, float(alloc)])
                    assignment_count += 1

            # All 12 months
            for m_idx, alloc_val in enumerate(alloc_values):
                if alloc_val is not None and alloc_val > 0:
                    con.execute("""
                        INSERT OR REPLACE INTO monthly_allocations
                        (employee_id, project_id, year, month, allocation, plan_version_id)
                        VALUES (?, ?, 2026, ?, ?, ?)
                    """, [workday_id, proj_slug, m_idx + 1, float(alloc_val), default_plan_version_id])
                    monthly_count += 1

        con.commit()
        print(f"  ✓ Inserted {assignment_count} active assignments (month: {MONTHS[CURRENT_MONTH_IDX]})")
        print(f"  ✓ Inserted {monthly_count} monthly allocation records (all 12 months)")
        print(f"  ✓ Inserted {placeholder_monthly_count} placeholder monthly allocation records")

        # ── Apply memory_facts overrides (e.g. resigned staff) ──
        resigned = con.execute("""
            SELECT subject FROM memory_facts
            WHERE category IN ('team_context', 'correction')
              AND (fact LIKE '%left the company%' OR fact LIKE '%inactive%' OR fact LIKE '%resigned%' OR fact LIKE '%离职%')
        """).fetchall()
        for (emp_id,) in resigned:
            resolved = con.execute(
                "SELECT id FROM employees WHERE id = ?",
                [emp_id],
            ).fetchone()
            if not resolved:
                resolved = con.execute(
                    """
                    SELECT employee_id AS id
                    FROM employee_external_ids
                    WHERE system_name = 'resource_portal'
                      AND lower(external_id) = lower(?)
                    LIMIT 1
                    """,
                    [emp_id],
                ).fetchone()
            target_id = resolved[0] if resolved else emp_id
            result = con.execute(
                "UPDATE employees SET status='inactive' WHERE id=? AND status='active'",
                [target_id],
            )
            if result.rowcount > 0:
                override_count += 1
        if override_count:
            con.commit()
            print(f"  ✓ Applied {override_count} memory_facts status override(s) (resigned/inactive staff)")

        con.close()
        con = None

        rows_changed = (
            len(employee_map)
            + len(placeholder_map)
            + ext_count
            + len(project_map)
            + 1
            + assignment_count
            + monthly_count
            + placeholder_monthly_count
            + override_count
        )
        repository.finish_sync_run(
            run_id,
            rows_in=len(records),
            rows_changed=rows_changed,
            notes="Distribution Excel import completed successfully.",
        )
        print()
        print("⚠️  Legacy import complete.")
        print(
            "   Current-state staffing product reads now use canonical publications,"
        )
        print(
            "   not direct assignment rows from this script."
        )
        print(
            "   Use `python3 scripts/import_team_project_capacity_workbook.py --file <workbook> --confirm`"
        )
        print("   for authoritative workbook onboarding.")
    except Exception as exc:
        if con is not None:
            con.close()
        repository.fail_sync_run(run_id, str(exc), rows_in=len(records))
        raise


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description='Import Distribution Excel into pm.db')
    parser.add_argument('--file', required=True, help='Path to Excel file')
    parser.add_argument('--dry-run', action='store_true')
    args = parser.parse_args()
    import_data(args.file, dry_run=args.dry_run)
