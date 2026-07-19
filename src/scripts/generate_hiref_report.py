"""
scripts/generate_hiref_report.py — Generate HIREF Status Report Excel from pm.db

Produces: data-feed/HIREF_Status_Report_<date>.xlsx
  Sheet 1: Staff HIREF Status
  Sheet 2: HIREF Allocation
  Sheet 3: Action Plan

Usage:
    python3 scripts/generate_hiref_report.py
"""

import sys
import sqlite3
from pathlib import Path
from datetime import date, datetime

sys.path.insert(0, str(Path(__file__).parent.parent))
from pm_agent.config import settings
from pm_agent.rules.hiref import project_alignment_status

try:
    from openpyxl import Workbook
    from openpyxl.styles import (PatternFill, Font, Alignment, Border, Side,
                                  GradientFill)
    from openpyxl.utils import get_column_letter
except ImportError:
    print("Run: pip3 install openpyxl")
    sys.exit(1)

TODAY = date.today()
GENERATED = TODAY.strftime('%Y-%m-%d')

# ── Colour palette ──────────────────────────────────────────────────────────
RED    = PatternFill("solid", fgColor="FF4444")
ORANGE = PatternFill("solid", fgColor="FF8C00")
YELLOW = PatternFill("solid", fgColor="FFD700")
GREEN  = PatternFill("solid", fgColor="70AD47")
BLUE   = PatternFill("solid", fgColor="1F6BB0")   # header
LBLUE  = PatternFill("solid", fgColor="BDD7EE")   # subheader
LGRAY  = PatternFill("solid", fgColor="F2F2F2")   # alt row
WHITE  = PatternFill("solid", fgColor="FFFFFF")

BOLD   = Font(bold=True)
WBOLD  = Font(bold=True, color="FFFFFF")
CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)
LEFT   = Alignment(horizontal='left', vertical='center', wrap_text=True)

def thin_border():
    s = Side(style='thin', color='CCCCCC')
    return Border(left=s, right=s, top=s, bottom=s)

def urgency_fill(days: int) -> PatternFill:
    if days <= 57:   return RED
    if days <= 118:  return ORANGE
    if days <= 210:  return YELLOW
    return GREEN

def urgency_label(days: int) -> str:
    if days <= 57:   return f'URGENT (<{days}d)'
    if days <= 118:  return f'Soon (<{days}d)'
    if days <= 210:  return f'Monitor (<{days}d)'
    return 'OK'

def effective_urgency_fill(days: int, has_next_hiref: bool) -> PatternFill:
    if has_next_hiref:
        return GREEN
    return urgency_fill(days)

def effective_urgency_label(days: int, has_next_hiref: bool) -> str:
    if has_next_hiref:
        return 'Next Confirmed'
    return urgency_label(days)

def format_period(start_date: str, end_date: str) -> str:
    if start_date and end_date:
        return f'{start_date} → {end_date}'
    if end_date:
        return f'? → {end_date}'
    return start_date or ''

def match_projects(hiref_project: str, actual_project: str) -> str:
    """Return the report label for the shared project-alignment rule."""
    status = project_alignment_status(
        hiref_project,
        actual_project_names=[actual_project] if actual_project else [],
    )
    return 'OK' if status == 'aligned' else 'MISMATCH'

def style_header_row(ws, row_num: int, cols: int, fill: PatternFill, font: Font):
    for c in range(1, cols + 1):
        cell = ws.cell(row=row_num, column=c)
        cell.fill = fill
        cell.font = font
        cell.alignment = CENTER
        cell.border = thin_border()

def autofit(ws, min_w=10, max_w=50):
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            try:
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
            except Exception:
                pass
        ws.column_dimensions[col_letter].width = min(max(max_len + 2, min_w), max_w)


# ── Load data from DB ────────────────────────────────────────────────────────

def load_data():
    con = sqlite3.connect(settings.database_path)
    con.row_factory = sqlite3.Row

    # STFTE staff with current project
    staff = con.execute("""
        SELECT e.name, e.role, e.resource_type, e.current_hiref, e.next_hiref,
               e.billing_end_date,
               ch.project AS current_hiref_project,
               ch.start_date AS current_hiref_start_date,
               ch.end_date AS current_hiref_end_date,
               nh.project AS next_hiref_project,
               nh.start_date AS next_hiref_start_date,
               nh.end_date AS next_hiref_end_date,
               COALESCE(GROUP_CONCAT(p.name, ' / '), '') as actual_projects
        FROM employees e
        LEFT JOIN hiref ch ON ch.id = e.current_hiref
        LEFT JOIN hiref nh ON nh.id = e.next_hiref
        LEFT JOIN assignments a ON a.employee_id = e.id
        LEFT JOIN projects p ON p.id = a.project_id
        WHERE e.resource_type = 'STFTE'
        GROUP BY e.id
        ORDER BY COALESCE(ch.end_date, e.billing_end_date), e.name
    """).fetchall()

    # HIREF contracts
    hiref = con.execute("""
        SELECT id, project, end_date, notes FROM hiref ORDER BY end_date
    """).fetchall()

    # Build hiref_id → project lookup
    hiref_lookup = {r['id']: r['project'] for r in hiref}

    con.close()
    return staff, hiref, hiref_lookup


# ── Sheet 1: Staff HIREF Status ──────────────────────────────────────────────

def build_staff_sheet(ws, staff, hiref_lookup):
    # Title
    ws.merge_cells('A1:L1')
    title = ws['A1']
    title.value = f'Yancey Team - Staff HIREF Status  |  Generated: {GENERATED}'
    title.fill = BLUE
    title.font = Font(bold=True, color='FFFFFF', size=13)
    title.alignment = CENTER
    ws.row_dimensions[1].height = 28

    # Headers
    headers = ['#', 'Name', 'Role', 'Type', 'Current HIREF', 'Current Period',
               'Next HIREF', 'Next Period', 'HIREF Project', 'Actual Project',
               'Match', 'Renewal Status']
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=2, column=c, value=h)
        cell.fill = LBLUE
        cell.font = BOLD
        cell.alignment = CENTER
        cell.border = thin_border()
    ws.row_dimensions[2].height = 20

    # Data
    for i, s in enumerate(staff):
        row = i + 3
        fill_row = LGRAY if i % 2 == 0 else WHITE

        current_start = s['current_hiref_start_date'] or ''
        end_date = s['current_hiref_end_date'] or s['billing_end_date'] or ''
        days = (datetime.strptime(end_date, '%Y-%m-%d').date() - TODAY).days if end_date else 999
        hiref_project = s['current_hiref_project'] or hiref_lookup.get(s['current_hiref'], '')
        current_period = format_period(current_start, end_date)
        next_start = s['next_hiref_start_date'] or ''
        next_end = s['next_hiref_end_date'] or ''
        next_period = format_period(next_start, next_end)
        actual = s['actual_projects'] or ''
        match = match_projects(hiref_project, actual) if hiref_project and actual else '?'
        has_next_hiref = bool(s['next_hiref'])
        renewal_status = f"Next ready: {s['next_hiref']}" if has_next_hiref else 'Next HIREF not assigned'

        values = [i + 1, s['name'], s['role'] or 'contractor', s['resource_type'],
                  s['current_hiref'], current_period, s['next_hiref'], next_period,
                  hiref_project, actual, match, renewal_status]

        for c, v in enumerate(values, 1):
            cell = ws.cell(row=row, column=c, value=v)
            cell.border = thin_border()
            cell.alignment = LEFT if c in (2, 6, 8, 9, 10, 12) else CENTER

            if c == 6:   # Current period
                cell.fill = effective_urgency_fill(days, has_next_hiref) if end_date else fill_row
                if end_date:
                    cell.font = Font(bold=True)
            elif c == 11:  # Match
                cell.fill = RED if match == 'MISMATCH' else GREEN
                cell.font = Font(bold=True)
            elif c == 12:  # Renewal status
                cell.fill = GREEN if s['next_hiref'] else (ORANGE if days <= 118 else fill_row)
            else:
                cell.fill = fill_row

    ws.freeze_panes = 'A3'
    autofit(ws)
    ws.column_dimensions['B'].width = 28
    ws.column_dimensions['F'].width = 24
    ws.column_dimensions['H'].width = 24
    ws.column_dimensions['I'].width = 42
    ws.column_dimensions['J'].width = 40
    ws.column_dimensions['L'].width = 26


# ── Sheet 2: HIREF Allocation ────────────────────────────────────────────────

def build_allocation_sheet(ws, staff, hiref, hiref_lookup):
    current_holder_map = {}
    next_holder_map = {}
    for s in staff:
        hid = s['current_hiref']
        if hid:
            current_holder_map.setdefault(hid, []).append({
                'name': s['name'],
                'has_next_hiref': bool(s['next_hiref']),
            })
        next_hid = s['next_hiref']
        if next_hid:
            next_holder_map.setdefault(next_hid, []).append(s['name'])

    ws.merge_cells('A1:I1')
    title = ws['A1']
    title.value = f'All HIREF Slots - Allocation Overview  |  Generated: {GENERATED}'
    title.fill = BLUE
    title.font = Font(bold=True, color='FFFFFF', size=13)
    title.alignment = CENTER
    ws.row_dimensions[1].height = 28

    headers = ['HIREF ID', 'Project', 'End Date', 'Days Left', 'Urgency',
               'Holder', 'Status', 'Action Required', 'Notes']
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=2, column=c, value=h)
        cell.fill = LBLUE
        cell.font = BOLD
        cell.alignment = CENTER
        cell.border = thin_border()
    ws.row_dimensions[2].height = 20

    row_num = 3
    for h_rec in hiref:
        hid = h_rec['id']
        end_date = h_rec['end_date'] or ''
        project = h_rec['project'] or ''
        notes = h_rec['notes'] or ''
        days = (datetime.strptime(end_date, '%Y-%m-%d').date() - TODAY).days if end_date else 999
        current_holders = current_holder_map.get(hid, [])
        next_holders = next_holder_map.get(hid, [])

        if current_holders:
            entries = []
            for holder in current_holders:
                action = 'Next HIREF confirmed' if holder['has_next_hiref'] else (
                    f'Renew before {end_date}' if end_date and days <= 118 else 'Monitor'
                )
                entries.append((holder['name'], 'Current', action, holder['has_next_hiref']))
        elif next_holders:
            action = 'Auto continue after current HIREF expiry'
            entries = [(holder, 'Next Planned', action, True) for holder in next_holders]
        else:
            action = 'Assign replacement slot' if 'replacement' in notes.lower() else 'Keep available'
            entries = [('', 'Available', action, False)]

        for holder, status, action, is_next_planned in entries:
            i_row = row_num - 3
            fill_row = LGRAY if i_row % 2 == 0 else WHITE
            values = [hid, project, end_date, days if end_date else '', effective_urgency_label(days, is_next_planned) if end_date else 'OK',
                      holder, status, action, notes]
            for c, v in enumerate(values, 1):
                cell = ws.cell(row=row_num, column=c, value=v)
                cell.border = thin_border()
                cell.alignment = LEFT if c in (2, 6, 8, 9) else CENTER
                if c == 4 and end_date:
                    cell.fill = effective_urgency_fill(days, is_next_planned)
                    cell.font = Font(bold=True)
                elif c == 5 and end_date:
                    cell.fill = effective_urgency_fill(days, is_next_planned)
                else:
                    cell.fill = fill_row
            row_num += 1

    ws.freeze_panes = 'A3'
    autofit(ws)
    ws.column_dimensions['B'].width = 44
    ws.column_dimensions['F'].width = 28
    ws.column_dimensions['H'].width = 28


# ── Sheet 3: Action Plan ──────────────────────────────────────────────────────

def build_action_sheet(ws, staff, hiref_lookup):
    ws.merge_cells('A1:F1')
    title = ws['A1']
    title.value = f'HIREF Action Plan  |  Generated: {GENERATED}'
    title.fill = BLUE
    title.font = Font(bold=True, color='FFFFFF', size=13)
    title.alignment = CENTER
    ws.row_dimensions[1].height = 28

    headers = ['Priority', 'Person', 'Current HIREF', 'Issue',
               'Recommended Action', 'Owner / Notes']
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=2, column=c, value=h)
        cell.fill = LBLUE
        cell.font = BOLD
        cell.alignment = CENTER
        cell.border = thin_border()

    actions = []
    for s in staff:
        end_date = s['current_hiref_end_date'] or s['billing_end_date'] or ''
        if not end_date:
            continue
        days = (datetime.strptime(end_date, '%Y-%m-%d').date() - TODAY).days
        hiref_project = s['current_hiref_project'] or hiref_lookup.get(s['current_hiref'], '')
        actual = s['actual_projects'] or ''
        match = match_projects(hiref_project, actual) if hiref_project and actual else '?'
        is_mismatch = (match == 'MISMATCH')
        next_hiref = s['next_hiref'] or ''
        next_end = s['next_hiref_end_date'] or ''

        if next_hiref:
            pri = '4 - OK (Next Confirmed)'
            fill = GREEN
        elif days <= 57:
            pri = '1 - URGENT Jul31'
            fill = RED
        elif days <= 118:
            pri = '2 - HIGH Sep30'
            fill = ORANGE
        elif days <= 210:
            pri = '3 - MEDIUM Dec31'
            fill = YELLOW
        else:
            pri = '4 - OK'
            fill = GREEN

        if is_mismatch:
            issue = f'MISMATCH: HIREF on {hiref_project}, working on {actual}'
            rec_action = f'Fix HIREF project → register under correct project'
            if days <= 57:
                pri = '0 - CRITICAL (Mismatch + Expiring)'
                fill = PatternFill("solid", fgColor="8B0000")
        elif next_hiref:
            issue = f'Current expires {end_date} ({days}d left)'
            rec_action = f'Next HIREF {next_hiref} ready' + (f' until {next_end}' if next_end else '')
        elif days <= 118:
            issue = f'Expires {end_date} ({days}d left)'
            rec_action = f'Assign next HIREF before {end_date}'
        else:
            issue = f'Expires {end_date}'
            rec_action = 'Monitor'

        actions.append((pri, s['name'], s['current_hiref'], issue, rec_action, '', fill, days, is_mismatch))

    # Sort: mismatches+urgent first, then by days
    actions.sort(key=lambda x: (0 if x[8] and x[7] <= 57 else (1 if x[7] <= 57 else (2 if x[7] <= 118 else 3)), x[7]))

    for i, (pri, name, hiref_id, issue, rec, notes, fill, days, _) in enumerate(actions):
        row = i + 3
        values = [pri, name, hiref_id, issue, rec, notes]
        for c, v in enumerate(values, 1):
            cell = ws.cell(row=row, column=c, value=v)
            cell.border = thin_border()
            cell.alignment = LEFT
            if c == 1:
                cell.fill = fill
                cell.font = Font(bold=True, color='FFFFFF' if days <= 118 else '000000')
            else:
                cell.fill = LGRAY if i % 2 == 0 else WHITE

    ws.freeze_panes = 'A3'
    autofit(ws)
    ws.column_dimensions['A'].width = 30
    ws.column_dimensions['B'].width = 28
    ws.column_dimensions['D'].width = 50
    ws.column_dimensions['E'].width = 45


# ── Main ──────────────────────────────────────────────────────────────────────

def main():
    staff, hiref, hiref_lookup = load_data()
    print(f"Loaded {len(staff)} STFTE staff, {len(hiref)} HIREF contracts")

    wb = Workbook()
    ws1 = wb.active
    ws1.title = 'Staff HIREF Status'
    ws2 = wb.create_sheet('HIREF Allocation')
    ws3 = wb.create_sheet('Action Plan')

    build_staff_sheet(ws1, staff, hiref_lookup)
    build_allocation_sheet(ws2, staff, hiref, hiref_lookup)
    build_action_sheet(ws3, staff, hiref_lookup)

    output = Path(settings.database_path).parent.parent / 'data-feed' / f'HIREF_Status_Report_{GENERATED}.xlsx'
    wb.save(str(output))
    print(f"✅ Report saved: {output}")

    # Summary
    from collections import Counter
    buckets = Counter()
    mismatches = []
    for s in staff:
        end_date = s['current_hiref_end_date'] or s['billing_end_date'] or ''
        if not end_date:
            continue
        days = (datetime.strptime(end_date, '%Y-%m-%d').date() - TODAY).days
        hiref_project = s['current_hiref_project'] or hiref_lookup.get(s['current_hiref'], '')
        actual = s['actual_projects'] or ''
        match = match_projects(hiref_project, actual) if hiref_project and actual else '?'
        if match == 'MISMATCH':
            mismatches.append(s['name'])
        if s['next_hiref']:
            buckets['SECURED (Next Confirmed)'] += 1
        elif days <= 57:   buckets['CRITICAL (Jul 31)'] += 1
        elif days <= 118: buckets['HIGH (Sep 30)'] += 1
        elif days <= 210: buckets['MEDIUM (Dec 31)'] += 1
        else:             buckets['OK (Apr 2027+)'] += 1

    print()
    print("STFTE Summary:")
    for label, cnt in buckets.items():
        print(f"  {label}: {cnt}")
    print(f"  Mismatches: {len(mismatches)} — {', '.join(mismatches)}")


if __name__ == '__main__':
    main()
