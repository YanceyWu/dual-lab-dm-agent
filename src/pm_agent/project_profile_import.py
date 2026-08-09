"""Project-profile workbook owner logic shared by onboarding and legacy wrapper."""

from __future__ import annotations

import hashlib
import json
import sqlite3
from pathlib import Path
from typing import Any

import openpyxl

from pm_agent.config import settings
from pm_agent.data_onboarding.source_export import (
    SourceExportError,
    failed_export,
    write_atomically,
)

WORKSHEET_NAME = "Project Profiles"
EXPECTED_HEADERS = (
    "Project Name",
    "Team Size",
    "Phase",
    "Phase Detail",
    "Priority",
    "Focus",
    "Objective",
    "Milestones",
    "Risks",
    "Stakeholders",
    "Special Rules",
    "Project ID",
)


def parse_milestones(text: object) -> list[dict[str, str]]:
    if not text or not str(text).strip():
        return []
    result: list[dict[str, str]] = []
    for line in str(text).strip().split("\n"):
        item = line.strip()
        if "|" not in item:
            continue
        parts = item.split("|", 1)
        result.append({"date": parts[0].strip(), "name": parts[1].strip()})
    return result


def parse_risks(text: object) -> list[dict[str, str]]:
    if not text or not str(text).strip():
        return []
    result: list[dict[str, str]] = []
    for line in str(text).strip().split("\n"):
        item = line.strip()
        if "|" in item:
            parts = item.split("|", 1)
            level = parts[1].strip().lower()
            if level not in {"high", "medium", "low"}:
                level = "medium"
            result.append({"risk": parts[0].strip(), "level": level})
        elif item:
            result.append({"risk": item, "level": "medium"})
    return result


def parse_stakeholders(text: object) -> list[str]:
    if not text or not str(text).strip():
        return []
    return [item.strip() for item in str(text).split(",") if item.strip()]


def safe(value: object, default: str = "") -> str:
    if value is None:
        return default
    text = str(value).strip()
    return text if text and text.lower() != "none" else default


def preview_project_profile_import(
    file_path: Path,
    *,
    db_path: str | Path | None = None,
) -> dict[str, Any]:
    plan = _build_import_plan(file_path, db_path=db_path)
    status = (
        "already_completed"
        if plan["blockers"] == []
        and int(plan["counts"].get("planned_profile_updates", 0)) == 0
        else "previewed"
    )
    return {
        "status": "rejected" if plan["blockers"] else status,
        "blockers": plan["blockers"],
        "warnings": plan["warnings"],
        "conflicts": [],
        "counts": plan["counts"],
        "payload": plan["payload"],
        "report": plan["report"],
    }


def apply_project_profile_import(
    file_path: Path,
    *,
    db_path: str | Path | None = None,
) -> dict[str, Any]:
    plan = _build_import_plan(file_path, db_path=db_path)
    if plan["blockers"]:
        blocker = plan["blockers"][0]
        raise ValueError(str(blocker["message"]))

    database = sqlite3.connect(_db_path(db_path))
    written_profiles = 0
    try:
        for row in plan["actionable_rows"]:
            database.execute(
                """
                INSERT INTO project_profiles
                    (project_id, phase, phase_detail, priority_tier, is_focus,
                     objective, milestones, key_risks, stakeholders, special_rules, updated_at)
                VALUES (?,?,?,?,?,?,?,?,?,?,datetime('now'))
                ON CONFLICT(project_id) DO UPDATE SET
                    phase=excluded.phase,
                    phase_detail=excluded.phase_detail,
                    priority_tier=excluded.priority_tier,
                    is_focus=excluded.is_focus,
                    objective=excluded.objective,
                    milestones=excluded.milestones,
                    key_risks=excluded.key_risks,
                    stakeholders=excluded.stakeholders,
                    special_rules=excluded.special_rules,
                    updated_at=excluded.updated_at
                """,
                [
                    row["project_id"],
                    row["phase"],
                    row["phase_detail"],
                    row["priority_tier"],
                    row["is_focus"],
                    row["objective"],
                    row["milestones_json"],
                    row["key_risks_json"],
                    row["stakeholders_json"],
                    row["special_rules"],
                ],
            )
            written_profiles += 1
        database.commit()
    except Exception:
        database.rollback()
        raise
    finally:
        database.close()

    report = dict(plan["report"])
    report["result"] = {
        "written_profiles": written_profiles,
        "skipped_phase_rows": plan["counts"]["skipped_phase_rows"],
        "missing_project_rows": plan["counts"]["missing_project_rows"],
    }
    return {
        "status": "completed",
        "blockers": [],
        "warnings": plan["warnings"],
        "conflicts": [],
        "counts": plan["counts"],
        "payload": plan["payload"],
        "report": report,
    }


def export_project_profile_workbook(
    output: str | Path,
    *,
    overwrite: bool = False,
    db_path: str | Path | None = None,
) -> dict[str, Any]:
    """Export only editable project-profile facts in the existing workbook shape."""

    source_type = "project-profile-workbook"
    connection = sqlite3.connect(_db_path(db_path))
    try:
        rows = connection.execute(
            """
            SELECT p.id, p.name, pp.phase, pp.phase_detail, pp.priority_tier,
                   pp.is_focus, pp.objective, pp.milestones, pp.key_risks,
                   pp.stakeholders, pp.special_rules
            FROM project_profiles pp
            JOIN projects p ON p.id = pp.project_id
            ORDER BY p.id
            """
        ).fetchall()
    except sqlite3.OperationalError:
        return failed_export(source_type, "PROJECT_PROFILE_EXPORT_SOURCE_UNAVAILABLE")
    finally:
        connection.close()

    try:
        expected_rows = [_exportable_profile_row(row) for row in rows]
    except ValueError:
        return failed_export(source_type, "PROJECT_PROFILE_EXPORT_SOURCE_UNREPRESENTABLE")

    def write(path: Path) -> None:
        workbook = openpyxl.Workbook()
        worksheet = workbook.active
        worksheet.title = WORKSHEET_NAME
        worksheet.append(["Project profiles"])
        worksheet.append(["Exported user-maintained source facts"])
        worksheet.append([None] * len(EXPECTED_HEADERS))
        worksheet.append(list(EXPECTED_HEADERS))
        for row in expected_rows:
            worksheet.append(
                [
                    row["project_name"],
                    "",
                    row["phase"],
                    row["phase_detail"],
                    row["priority_tier"],
                    "Y" if row["is_focus"] else "",
                    row["objective"],
                    _milestones_text(row["milestones"]),
                    _risks_text(row["key_risks"]),
                    _stakeholders_text(row["stakeholders"]),
                    row["special_rules"],
                    row["project_id"],
                ]
            )
        workbook.save(path)
        preview = preview_project_profile_import(path, db_path=db_path)
        plan = _build_import_plan(path, db_path=db_path)
        actual_rows = [
            _profile_facts_from_candidate(candidate)
            for candidate in plan["actionable_rows"]
        ]
        expected_facts = [_profile_facts_from_export(row) for row in expected_rows]
        if (
            preview["status"] == "rejected"
            or len(actual_rows) != len(expected_facts)
            or actual_rows != expected_facts
        ):
            raise SourceExportError("PROJECT_PROFILE_EXPORT_SOURCE_UNREPRESENTABLE")

    try:
        target, digest = write_atomically(output, overwrite=overwrite, write=write)
    except SourceExportError as exc:
        return failed_export(source_type, str(exc))
    return {
        "status": "exported",
        "source_type": source_type,
        "output_path": str(target),
        "sha256": digest,
        "counts": {"project_profiles": len(rows)},
        "warnings": [],
        "source_identity": {"kind": "local_sqlite", "source_type": source_type},
    }


def _build_import_plan(
    file_path: Path,
    *,
    db_path: str | Path | None = None,
) -> dict[str, Any]:
    workbook = openpyxl.load_workbook(file_path, data_only=True)
    if WORKSHEET_NAME not in workbook.sheetnames:
        return {
            "blockers": [
                {
                    "severity": "blocker",
                    "code": "PROJECT_PROFILE_WORKBOOK_SHEET_MISSING",
                    "message": f"Workbook must contain worksheet '{WORKSHEET_NAME}'.",
                    "location": "source_locator",
                }
            ],
            "warnings": [],
            "counts": {
                "source_rows": 0,
                "actionable_profiles": 0,
                "planned_profile_updates": 0,
                "unchanged_profiles": 0,
                "skipped_phase_rows": 0,
                "missing_project_rows": 0,
            },
            "payload": {"publication_id": _publication_id([]), "preview_rows": []},
            "report": {"coverage": {"state": "not_available"}},
            "actionable_rows": [],
        }

    worksheet = workbook[WORKSHEET_NAME]
    actionable_rows: list[dict[str, Any]] = []
    preview_rows: list[dict[str, Any]] = []
    warnings: list[dict[str, Any]] = []
    counts = {
        "source_rows": 0,
        "actionable_profiles": 0,
        "planned_profile_updates": 0,
        "unchanged_profiles": 0,
        "skipped_phase_rows": 0,
        "missing_project_rows": 0,
    }

    database = sqlite3.connect(_db_path(db_path))
    database.row_factory = sqlite3.Row
    try:
        known_projects = {
            str(row[0])
            for row in database.execute("SELECT id FROM projects").fetchall()
            if row[0] is not None
        }
        existing_profiles = {
            str(row["project_id"]): dict(row)
            for row in database.execute("SELECT * FROM project_profiles").fetchall()
        }
    finally:
        database.close()

    for row_number, row in enumerate(
        worksheet.iter_rows(min_row=4, values_only=True),
        start=4,
    ):
        padded = list(row[: len(EXPECTED_HEADERS)]) + [None] * max(0, len(EXPECTED_HEADERS) - len(row))
        (
            project_name,
            _team_size,
            phase,
            phase_detail,
            priority,
            focus,
            objective,
            milestones_raw,
            risks_raw,
            stakeholders_raw,
            special_rules,
            project_id,
        ) = padded[: len(EXPECTED_HEADERS)]

        project_id_text = safe(project_id)
        if not project_id_text:
            continue
        if project_id_text.lower() == "project id":
            continue

        counts["source_rows"] += 1
        label = safe(project_name, project_id_text)
        phase_text = safe(phase)
        if not phase_text or phase_text == "TBC":
            counts["skipped_phase_rows"] += 1
            preview_rows.append(
                {
                    "row": row_number,
                    "project_id": project_id_text,
                    "project_name": label,
                    "action": "skipped",
                    "reason": "phase_not_filled",
                }
            )
            continue
        if project_id_text not in known_projects:
            counts["missing_project_rows"] += 1
            preview_rows.append(
                {
                    "row": row_number,
                    "project_id": project_id_text,
                    "project_name": label,
                    "action": "error",
                    "reason": "project_not_found",
                }
            )
            continue

        priority_tier = 2
        try:
            priority_tier = int(priority)
            if priority_tier not in (1, 2, 3):
                priority_tier = 2
        except (TypeError, ValueError):
            priority_tier = 2

        candidate = {
            "project_id": project_id_text,
            "project_name": label,
            "phase": phase_text,
            "phase_detail": safe(phase_detail),
            "priority_tier": priority_tier,
            "is_focus": 1 if safe(focus).upper() == "Y" else 0,
            "objective": safe(objective),
            "milestones": parse_milestones(milestones_raw),
            "key_risks": parse_risks(risks_raw),
            "stakeholders": parse_stakeholders(stakeholders_raw),
            "special_rules": safe(special_rules),
        }
        candidate["milestones_json"] = json.dumps(candidate["milestones"], ensure_ascii=False)
        candidate["key_risks_json"] = json.dumps(candidate["key_risks"], ensure_ascii=False)
        candidate["stakeholders_json"] = json.dumps(
            candidate["stakeholders"], ensure_ascii=False
        )

        counts["actionable_profiles"] += 1
        existing = existing_profiles.get(project_id_text)
        changed = not _project_profile_matches(existing, candidate)
        if changed:
            counts["planned_profile_updates"] += 1
        else:
            counts["unchanged_profiles"] += 1
        preview_rows.append(
            {
                "row": row_number,
                "project_id": project_id_text,
                "project_name": label,
                "action": "update" if changed else "unchanged",
                "phase": candidate["phase"],
                "priority_tier": candidate["priority_tier"],
                "is_focus": candidate["is_focus"],
            }
        )
        actionable_rows.append(candidate)

    if counts["skipped_phase_rows"]:
        warnings.append(
            {
                "severity": "warning",
                "code": "PROJECT_PROFILE_PHASE_NOT_FILLED_SKIPPED",
                "message": (
                    f"{counts['skipped_phase_rows']} row(s) were skipped because phase is blank or TBC."
                ),
                "location": "source",
            }
        )
    if counts["missing_project_rows"]:
        warnings.append(
            {
                "severity": "warning",
                "code": "PROJECT_PROFILE_PROJECT_NOT_FOUND_SKIPPED",
                "message": (
                    f"{counts['missing_project_rows']} row(s) reference projects that do not exist locally and will be ignored."
                ),
                "location": "source",
            }
        )

    payload_rows = [
        {
            "project_id": row["project_id"],
            "phase": row["phase"],
            "phase_detail": row["phase_detail"],
            "priority_tier": row["priority_tier"],
            "is_focus": row["is_focus"],
            "objective": row["objective"],
            "milestones": row["milestones"],
            "key_risks": row["key_risks"],
            "stakeholders": row["stakeholders"],
            "special_rules": row["special_rules"],
        }
        for row in actionable_rows
    ]
    return {
        "blockers": [],
        "warnings": warnings,
        "counts": counts,
        "payload": {
            "publication_id": _publication_id(payload_rows),
            "profile_fingerprint": _fingerprint(payload_rows),
            "project_ids": [row["project_id"] for row in actionable_rows],
            "preview_rows": preview_rows,
        },
        "report": {
            "coverage": {
                "state": "complete",
                "source_rows": counts["source_rows"],
                "actionable_profiles": counts["actionable_profiles"],
                "planned_profile_updates": counts["planned_profile_updates"],
                "unchanged_profiles": counts["unchanged_profiles"],
                "skipped_phase_rows": counts["skipped_phase_rows"],
                "missing_project_rows": counts["missing_project_rows"],
            },
            "worksheet": WORKSHEET_NAME,
            "preserved_semantics": {
                "phase_required_for_update": True,
                "missing_project_rows_non_blocking": True,
                "idempotent_replay_short_circuits_unchanged_rows": True,
            },
        },
        "actionable_rows": actionable_rows,
    }


def _project_profile_matches(existing: dict[str, Any] | None, candidate: dict[str, Any]) -> bool:
    if existing is None:
        return False
    return {
        "phase": safe(existing.get("phase")),
        "phase_detail": safe(existing.get("phase_detail")),
        "priority_tier": int(existing.get("priority_tier") or 2),
        "is_focus": int(existing.get("is_focus") or 0),
        "objective": safe(existing.get("objective")),
        "milestones": _json_or_empty(existing.get("milestones")),
        "key_risks": _json_or_empty(existing.get("key_risks")),
        "stakeholders": _json_or_empty(existing.get("stakeholders")),
        "special_rules": safe(existing.get("special_rules")),
    } == {
        "phase": candidate["phase"],
        "phase_detail": candidate["phase_detail"],
        "priority_tier": candidate["priority_tier"],
        "is_focus": candidate["is_focus"],
        "objective": candidate["objective"],
        "milestones": candidate["milestones"],
        "key_risks": candidate["key_risks"],
        "stakeholders": candidate["stakeholders"],
        "special_rules": candidate["special_rules"],
    }


def _json_or_empty(value: object) -> object:
    if not value:
        return []
    try:
        return json.loads(str(value))
    except json.JSONDecodeError:
        return value


def _json_list(value: object) -> list[object]:
    decoded = _json_or_empty(value)
    return decoded if isinstance(decoded, list) else []


def _milestones_text(value: object) -> str:
    return "\n".join(
        f"{item.get('date', '')} | {item.get('name', '')}".strip()
        for item in _json_list(value)
        if isinstance(item, dict)
    )


def _risks_text(value: object) -> str:
    return "\n".join(
        f"{item.get('risk', '')} | {item.get('level', 'medium')}".strip()
        for item in _json_list(value)
        if isinstance(item, dict)
    )


def _stakeholders_text(value: object) -> str:
    return ", ".join(str(item) for item in _json_list(value) if str(item).strip())


def _exportable_profile_row(row: tuple[Any, ...]) -> dict[str, Any]:
    project_id, project_name, phase, phase_detail, priority_tier, is_focus, objective, milestones, risks, stakeholders, special_rules = row
    if (
        not isinstance(project_id, str)
        or not project_id.strip()
        or (project_name is not None and not _canonical_text(project_name))
        or not _canonical_text(phase)
        or not isinstance(priority_tier, int)
        or isinstance(priority_tier, bool)
        or priority_tier not in {1, 2, 3}
        or not isinstance(is_focus, int)
        or isinstance(is_focus, bool)
        or is_focus not in {0, 1}
    ):
        raise ValueError("unrepresentable project profile")
    return {
        "project_id": project_id,
        "project_name": project_name or project_id,
        "phase": str(phase),
        "phase_detail": _optional_canonical_text(phase_detail),
        "priority_tier": priority_tier,
        "is_focus": is_focus,
        "objective": _optional_canonical_text(objective),
        "milestones": _strict_json_list(milestones, {"date", "name"}),
        "key_risks": _strict_json_list(risks, {"risk", "level"}),
        "stakeholders": _strict_stakeholders(stakeholders),
        "special_rules": _optional_canonical_text(special_rules),
    }


def _canonical_text(value: object) -> bool:
    return isinstance(value, str) and bool(value.strip()) and value == value.strip()


def _optional_canonical_text(value: object) -> str:
    if value is None:
        return ""
    if not isinstance(value, str) or value != value.strip() or value.lower() == "none":
        raise ValueError("unrepresentable text")
    return value


def _strict_json_list(value: object, required_keys: set[str]) -> list[dict[str, str]]:
    if value is None or value == "":
        return []
    if not isinstance(value, str):
        raise ValueError("unrepresentable JSON")
    try:
        decoded = json.loads(value)
    except json.JSONDecodeError as exc:
        raise ValueError("malformed JSON") from exc
    if not isinstance(decoded, list):
        raise ValueError("malformed JSON")
    result: list[dict[str, str]] = []
    for item in decoded:
        if (
            not isinstance(item, dict)
            or set(item) != required_keys
            or any(not _canonical_text(item[key]) for key in required_keys)
            or any("|" in item[key] or "\n" in item[key] for key in required_keys)
        ):
            raise ValueError("unrepresentable JSON")
        if required_keys == {"risk", "level"} and item["level"] not in {"high", "medium", "low"}:
            raise ValueError("unrepresentable risk")
        result.append({key: item[key] for key in required_keys})
    return result


def _strict_stakeholders(value: object) -> list[str]:
    if value is None or value == "":
        return []
    if not isinstance(value, str):
        raise ValueError("unrepresentable stakeholders")
    try:
        decoded = json.loads(value)
    except json.JSONDecodeError as exc:
        raise ValueError("malformed stakeholders") from exc
    if not isinstance(decoded, list) or any(
        not _canonical_text(item) or "," in item for item in decoded
    ):
        raise ValueError("unrepresentable stakeholders")
    return list(decoded)


def _profile_facts_from_export(row: dict[str, Any]) -> dict[str, Any]:
    return {
        key: row[key]
        for key in (
            "project_id", "phase", "phase_detail", "priority_tier", "is_focus",
            "objective", "milestones", "key_risks", "stakeholders", "special_rules",
        )
    }


def _profile_facts_from_candidate(candidate: dict[str, Any]) -> dict[str, Any]:
    return {
        "project_id": candidate["project_id"],
        "phase": candidate["phase"],
        "phase_detail": candidate["phase_detail"],
        "priority_tier": candidate["priority_tier"],
        "is_focus": candidate["is_focus"],
        "objective": candidate["objective"],
        "milestones": candidate["milestones"],
        "key_risks": candidate["key_risks"],
        "stakeholders": candidate["stakeholders"],
        "special_rules": candidate["special_rules"],
    }


def _db_path(db_path: str | Path | None) -> Path:
    return Path(db_path or settings.database_path)


def _fingerprint(rows: list[dict[str, Any]]) -> str:
    digest = hashlib.sha256()
    digest.update(
        json.dumps(rows, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode(
            "utf-8"
        )
    )
    return digest.hexdigest()


def _publication_id(rows: list[dict[str, Any]]) -> str:
    return f"project-profile-workbook:{_fingerprint(rows)}"
