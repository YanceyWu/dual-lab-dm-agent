"""Workbook parsing for Team/Project + Capacity onboarding."""

from __future__ import annotations

from dataclasses import replace
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from pm_agent.workbook_onboarding.models import (
    ParsedWorkbook,
    ValidationIssue,
    WorkbookAllocationRow,
    WorkbookCapacityRow,
    WorkbookHirefMemberRow,
    WorkbookHirefPlaceholderAllocationRow,
    WorkbookHirefPlaceholderRow,
    WorkbookHirefSlotRow,
    WorkbookMemberRow,
    WorkbookProjectRow,
    WorkbookSetupRow,
)
from pm_agent.workbook_onboarding.presets import (
    DEFAULT_WORKBOOK_MAPPING_PRESET_ID,
    WorkbookHeaderResolution,
    WorkbookPreset,
    WorkbookPresetField,
    WorkbookPresetResolution,
    WorkbookPresetSection,
    WorkbookSectionResolution,
    get_workbook_preset,
)

_DEFAULT_PRESET = get_workbook_preset(DEFAULT_WORKBOOK_MAPPING_PRESET_ID)
EXPECTED_SHEETS = tuple(
    section.primary_sheet_name for section in _DEFAULT_PRESET.sections
)
EXPECTED_HEADERS = {
    section.primary_sheet_name: [field.header_name for field in section.fields]
    for section in _DEFAULT_PRESET.sections
}


class WorkbookParseError(ValueError):
    """Raised when the workbook structure does not match the selected preset."""

    def __init__(
        self,
        code: str,
        message: str,
        *,
        location: str = "workbook",
        preset_resolution: WorkbookPresetResolution | None = None,
    ) -> None:
        self.code = code
        self.message = message
        self.location = location
        self.preset_resolution = preset_resolution
        super().__init__(self._render())

    def _render(self) -> str:
        if self.location == "workbook":
            return self.code
        return f"{self.code}:{self.location}"


def parse_workbook(
    path: str | Path,
    *,
    mapping_preset_id: str = DEFAULT_WORKBOOK_MAPPING_PRESET_ID,
) -> ParsedWorkbook:
    preset = get_workbook_preset(mapping_preset_id)
    workbook = load_workbook(Path(path), data_only=True)
    matched_sections, missing_optional = _resolve_sections(
        tuple(workbook.sheetnames),
        preset,
    )
    section_resolutions = [resolution for _, resolution in matched_sections]
    contract_warnings = _section_contract_warnings(section_resolutions, missing_optional)
    rows_by_section: dict[str, list[tuple[int, list[Any]]]] = {}

    for index, (section, section_resolution) in enumerate(matched_sections):
        if section_resolution.resolved_sheet_name is None:
            rows_by_section[section.section_key] = []
            continue
        try:
            header_matches, rows = _read_sheet(
                workbook[section_resolution.resolved_sheet_name],
                section=section,
                preset=preset,
                section_resolutions=section_resolutions,
                missing_optional=missing_optional,
            )
        except WorkbookParseError as exc:
            if exc.preset_resolution is None:
                exc.preset_resolution = _build_resolution(
                    preset,
                    tuple(section_resolutions),
                    missing_optional=missing_optional,
                    missing_required=(),
                )
            raise
        section_resolutions[index] = replace(
            section_resolution,
            header_matches=header_matches,
        )
        rows_by_section[section.section_key] = rows
        contract_warnings.extend(
            _header_contract_warnings(
                header_matches,
                section=section,
                sheet_name=section_resolution.resolved_sheet_name,
            )
        )

    preset_resolution = _build_resolution(
        preset,
        tuple(section_resolutions),
        missing_optional=missing_optional,
        missing_required=(),
    )
    return ParsedWorkbook(
        setup_rows=[
            WorkbookSetupRow(
                row_number=row_number,
                plan_version_name=_as_text(values[0]),
                as_of_date=_as_date_text(values[1]),
                start_month=_as_month_text(values[2]),
                end_month=_as_month_text(values[3]),
            )
            for row_number, values in rows_by_section["setup"]
        ],
        members=[
            WorkbookMemberRow(
                row_number=row_number,
                member_key=_as_text(values[0]),
                display_name=_as_text(values[1]),
                fte_type=_as_text(values[2]),
                status=_as_text(values[3]),
                current_hiref_id=_as_text(values[4]),
                hiref_end_date=_as_date_text(values[5]),
                role=_as_text(values[6]),
                level=_as_int(values[7]),
                effective_start=_as_date_text(values[8]),
                effective_end=_as_date_text(values[9]),
            )
            for row_number, values in rows_by_section["members"]
        ],
        projects=[
            WorkbookProjectRow(
                row_number=row_number,
                project_key=_as_text(values[0]),
                display_name=_as_text(values[1]),
                status=_as_text(values[2]),
                priority=_as_int(values[3]),
                start_date=_as_date_text(values[4]),
                target_end=_as_date_text(values[5]),
            )
            for row_number, values in rows_by_section["projects"]
        ],
        allocations=[
            WorkbookAllocationRow(
                row_number=row_number,
                member_key=_as_text(values[0]),
                project_key=_as_text(values[1]),
                month=_as_month_text(values[2]),
                allocation=_as_float(values[3]),
            )
            for row_number, values in rows_by_section["allocations"]
        ],
        capacity_rows=[
            WorkbookCapacityRow(
                row_number=row_number,
                member_key=_as_text(values[0]),
                month=_as_month_text(values[1]),
                leave_fraction=_as_float(values[2]),
                bau_fraction=_as_float(values[3]),
                non_project_fraction=_as_float(values[4]),
            )
            for row_number, values in rows_by_section["capacity"]
            if not _is_capacity_note_row(values)
        ],
        hiref_members=[
            WorkbookHirefMemberRow(
                row_number=row_number,
                member_key=_as_text(values[0]),
                next_hiref_id=_as_text(values[1]),
            )
            for row_number, values in rows_by_section["hiref_members"]
        ],
        hiref_slots=[
            WorkbookHirefSlotRow(
                row_number=row_number,
                hiref_id=_as_text(values[0]),
                project=_as_text(values[1]),
                request_type=_as_text(values[2]),
                start_date=_as_date_text(values[3]),
                end_date=_as_date_text(values[4]),
                notes=_as_text(values[5]),
            )
            for row_number, values in rows_by_section["hiref_slots"]
        ],
        hiref_placeholders=[
            WorkbookHirefPlaceholderRow(
                row_number=row_number,
                placeholder_id=_as_text(values[0]),
                display_name=_as_text(values[1]),
                hiref_id=_as_text(values[2]),
                linked_member_key=_as_text(values[3]),
                resource_type=_as_text(values[4]),
                status=_as_text(values[5]),
                notes=_as_text(values[6]),
            )
            for row_number, values in rows_by_section["hiref_placeholders"]
        ],
        hiref_placeholder_allocations=[
            WorkbookHirefPlaceholderAllocationRow(
                row_number=row_number,
                placeholder_id=_as_text(values[0]),
                project_key=_as_text(values[1]),
                month=_as_month_text(values[2]),
                allocation=_as_float(values[3]),
            )
            for row_number, values in rows_by_section["hiref_placeholder_allocations"]
        ],
        preset_resolution=preset_resolution,
        contract_warnings=tuple(contract_warnings),
    )


def _resolve_sections(
    sheet_names: tuple[str, ...],
    preset: WorkbookPreset,
) -> tuple[
    list[tuple[WorkbookPresetSection, WorkbookSectionResolution]],
    tuple[str, ...],
]:
    matched_sections: list[tuple[WorkbookPresetSection, WorkbookSectionResolution]] = []
    used_sheet_names: set[str] = set()
    missing_optional: list[str] = []
    missing_required: list[str] = []

    for section in preset.sections:
        matches = [name for name in sheet_names if name in section.accepted_sheet_names]
        if len(matches) > 1:
            partial_resolution = _build_resolution(
                preset,
                tuple(resolution for _, resolution in matched_sections)
                + (
                    WorkbookSectionResolution(
                        section_key=section.section_key,
                        display_name=section.display_name,
                        required=section.required,
                        expected_sheet_name=section.primary_sheet_name,
                        resolved_sheet_name=None,
                        matched_via_alias=False,
                    ),
                ),
                missing_optional=tuple(missing_optional),
                missing_required=tuple(missing_required),
            )
            raise WorkbookParseError(
                "WORKBOOK_SECTION_AMBIGUOUS",
                (
                    f"{section.display_name} matched more than one allowed sheet name: "
                    f"{', '.join(matches)}."
                ),
                location=section.display_name,
                preset_resolution=partial_resolution,
            )
        if not matches:
            resolution = WorkbookSectionResolution(
                section_key=section.section_key,
                display_name=section.display_name,
                required=section.required,
                expected_sheet_name=section.primary_sheet_name,
                resolved_sheet_name=None,
                matched_via_alias=False,
            )
            matched_sections.append((section, resolution))
            if section.required:
                missing_required.append(section.section_key)
            else:
                missing_optional.append(section.section_key)
            continue
        resolved_sheet_name = matches[0]
        used_sheet_names.add(resolved_sheet_name)
        matched_sections.append(
            (
                section,
                WorkbookSectionResolution(
                    section_key=section.section_key,
                    display_name=section.display_name,
                    required=section.required,
                    expected_sheet_name=section.primary_sheet_name,
                    resolved_sheet_name=resolved_sheet_name,
                    matched_via_alias=resolved_sheet_name != section.primary_sheet_name,
                ),
            )
        )

    unexpected = tuple(name for name in sheet_names if name not in used_sheet_names)
    if missing_required or unexpected:
        resolution = _build_resolution(
            preset,
            tuple(section_resolution for _, section_resolution in matched_sections),
            missing_optional=tuple(missing_optional),
            missing_required=tuple(missing_required),
        )
        details: list[str] = []
        if missing_required:
            details.append(f"missing required sections: {', '.join(missing_required)}")
        if unexpected:
            details.append(f"unexpected sheets: {', '.join(unexpected)}")
        raise WorkbookParseError(
            "WORKBOOK_SHEETS_INVALID",
            "Workbook sheets do not satisfy the selected preset (" + "; ".join(details) + ").",
            preset_resolution=resolution,
        )

    return matched_sections, tuple(missing_optional)


def _read_sheet(
    worksheet: Any,
    *,
    section: WorkbookPresetSection,
    preset: WorkbookPreset,
    section_resolutions: list[WorkbookSectionResolution],
    missing_optional: tuple[str, ...],
) -> tuple[tuple[WorkbookHeaderResolution, ...], list[tuple[int, list[Any]]]]:
    rows = list(worksheet.iter_rows(values_only=True))
    if not rows:
        raise WorkbookParseError(
            "WORKBOOK_HEADER_MISSING",
            f"{worksheet.title} is missing a header row.",
            location=worksheet.title,
            preset_resolution=_build_resolution(
                preset,
                tuple(section_resolutions),
                missing_optional=missing_optional,
                missing_required=(),
            ),
        )

    header_row = list(rows[0])
    header_texts = [_as_header_text(value) for value in header_row]
    field_matches = [_matching_header_indexes(header_texts, field) for field in section.fields]
    for field, indexes in zip(section.fields, field_matches, strict=True):
        if len(indexes) > 1:
            raise WorkbookParseError(
                "WORKBOOK_HEADER_AMBIGUOUS",
                (
                    f"{worksheet.title} contains more than one allowed header for "
                    f"{field.field_key}."
                ),
                location=f"{worksheet.title}:{field.field_key}",
                preset_resolution=_build_resolution(
                    preset,
                    tuple(section_resolutions),
                    missing_optional=missing_optional,
                    missing_required=(),
                ),
            )

    matched_indexes = [indexes[0] for indexes in field_matches if indexes]
    if (
        len(matched_indexes) != len(section.fields)
        or matched_indexes != list(range(len(section.fields)))
        or any(not _is_blank(value) for value in header_row[len(section.fields) :])
    ):
        raise WorkbookParseError(
            "WORKBOOK_HEADERS_INVALID",
            f"{worksheet.title} headers do not match the selected preset.",
            location=worksheet.title,
            preset_resolution=_build_resolution(
                preset,
                tuple(section_resolutions),
                missing_optional=missing_optional,
                missing_required=(),
            ),
        )

    header_matches = tuple(
        WorkbookHeaderResolution(
            field_key=field.field_key,
            expected_header_name=field.header_name,
            resolved_header_name=header_texts[indexes[0]],
            matched_via_alias=header_texts[indexes[0]] != field.header_name,
        )
        for field, indexes in zip(section.fields, field_matches, strict=True)
    )
    enriched_section_resolutions = _with_section_header_matches(
        section_resolutions,
        section_key=section.section_key,
        header_matches=header_matches,
    )
    results: list[tuple[int, list[Any]]] = []
    field_count = len(section.fields)
    for row_number, values in enumerate(rows[1:], start=2):
        row_values = list(values)
        if any(not _is_blank(value) for value in row_values[field_count:]):
            raise WorkbookParseError(
                "WORKBOOK_ROW_WIDTH_INVALID",
                f"{worksheet.title} row {row_number} contains unexpected extra values.",
                location=f"{worksheet.title}:{row_number}",
                preset_resolution=_build_resolution(
                    preset,
                    enriched_section_resolutions,
                    missing_optional=missing_optional,
                    missing_required=(),
                ),
            )
        current = row_values[:field_count]
        if len(current) < field_count:
            current.extend([None] * (field_count - len(current)))
        if all(_is_blank(value) for value in current):
            continue
        results.append((row_number, current))
    return header_matches, results


def _matching_header_indexes(
    header_texts: list[str],
    field: WorkbookPresetField,
) -> list[int]:
    return [
        index
        for index, text in enumerate(header_texts)
        if text in field.accepted_header_names
    ]


def _build_resolution(
    preset: WorkbookPreset,
    sections: tuple[WorkbookSectionResolution, ...],
    *,
    missing_optional: tuple[str, ...],
    missing_required: tuple[str, ...],
) -> WorkbookPresetResolution:
    return WorkbookPresetResolution(
        mapping_preset_id=preset.mapping_preset_id,
        display_name=preset.display_name,
        status=preset.status,
        matched_via_alias=any(
            section.matched_via_alias
            or any(match.matched_via_alias for match in section.header_matches)
            for section in sections
        ),
        sections=sections,
        missing_optional_sections=missing_optional,
        missing_required_sections=missing_required,
    )


def _with_section_header_matches(
    section_resolutions: list[WorkbookSectionResolution],
    *,
    section_key: str,
    header_matches: tuple[WorkbookHeaderResolution, ...],
) -> tuple[WorkbookSectionResolution, ...]:
    return tuple(
        replace(section_resolution, header_matches=header_matches)
        if section_resolution.section_key == section_key
        else section_resolution
        for section_resolution in section_resolutions
    )


def _section_contract_warnings(
    section_resolutions: list[WorkbookSectionResolution],
    missing_optional: tuple[str, ...],
) -> list[ValidationIssue]:
    warnings: list[ValidationIssue] = []
    optional_missing = set(missing_optional)
    for section in section_resolutions:
        if section.section_key in optional_missing:
            warnings.append(
                ValidationIssue(
                    "warning",
                    "WORKBOOK_OPTIONAL_SECTION_MISSING",
                    f"{section.display_name} is not present and will be treated as missing optional source coverage.",
                    section.display_name,
                )
            )
        elif section.matched_via_alias and section.resolved_sheet_name is not None:
            warnings.append(
                ValidationIssue(
                    "warning",
                    "WORKBOOK_SHEET_ALIAS_USED",
                    f"{section.display_name} matched allowed sheet alias {section.resolved_sheet_name!r}.",
                    section.resolved_sheet_name,
                )
            )
    return warnings


def _header_contract_warnings(
    header_matches: tuple[WorkbookHeaderResolution, ...],
    *,
    section: WorkbookPresetSection,
    sheet_name: str,
) -> list[ValidationIssue]:
    warnings: list[ValidationIssue] = []
    for match in header_matches:
        if match.matched_via_alias:
            warnings.append(
                ValidationIssue(
                    "warning",
                    "WORKBOOK_HEADER_ALIAS_USED",
                    (
                        f"{section.display_name} field {match.field_key} matched allowed "
                        f"header alias {match.resolved_header_name!r}."
                    ),
                    f"{sheet_name}!1",
                )
            )
    return warnings


def _is_blank(value: object) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _as_header_text(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _as_text(value: object) -> str | None:
    if _is_blank(value):
        return None
    if isinstance(value, (date, datetime)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    return str(value).strip()


def _as_date_text(value: object) -> str | None:
    if _is_blank(value):
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


def _as_month_text(value: object) -> str | None:
    if _is_blank(value):
        return None
    if isinstance(value, datetime):
        return value.strftime("%Y-%m")
    if isinstance(value, date):
        return value.strftime("%Y-%m")
    return str(value).strip()


def _as_int(value: object) -> int | None:
    if _is_blank(value):
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return None


def _as_float(value: object) -> float | None:
    if _is_blank(value):
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return float(str(value).strip())
    except ValueError:
        return None


def _is_capacity_note_row(values: list[Any]) -> bool:
    first = _as_text(values[0])
    if not first or not first.lower().startswith("note:"):
        return False
    return all(_is_blank(value) for value in values[1:])
